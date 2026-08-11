"""MarkdownFactory 1.4 — every document becomes Markdown (or Word).

1.4: a PDF engine for the Documents tab. Firecrawl's pdf-inspector reads a PDF
     the way a person does — it works out what kind of file it is before trying
     to read it, so a scan is turned away with a sentence instead of quietly
     producing an empty .md, and a PDF that does have text is read by a
     position-aware extractor that keeps the headings, tables, columns and
     links markitdown flattens away. Everything that is not a PDF takes exactly
     the road it took in 1.3.

1.3: To Word tab. Markdown is the right output for a knowledge base and the
     wrong one for anybody who has to hand the file back to a colleague; this
     tab is the way back out, from a PDF or from Markdown itself.

One window, four tabs (MangaFactory-style):

Tab 1 — Documents:      Word / PowerPoint / Excel / HTML / born-digital PDFs and
                        more, converted in-process. PDFs go through Firecrawl's
                        pdf-inspector (1.4), which classifies before it extracts
                        and sends scans to the Scans & Photos tab; every other
                        format goes through Microsoft's markitdown, which is
                        also the fallback if the PDF engine isn't aboard.
                        Bundled into the exe; works instantly on any PC.
Tab 2 — Scans & Photos: Scanned PDFs, photos of documents, screenshots — read by
                        Baidu's Unlimited-OCR vision model running locally on an
                        NVIDIA GPU. Too big to bundle (CUDA torch is ~4 GB), so
                        this tab manages its own runtime: one-click setup
                        downloads an embedded Python + the libraries into
                        %LOCALAPPDATA%\\MarkdownFactory, then runs the proven
                        Unlimited-OCR GUI server as a background worker and
                        embeds its page. No NVIDIA card -> the tab explains
                        that in plain English instead of failing.
Tab 3 — PDF Tools:      Merge / split / compress PDFs (absorbed from the PDF
                        Tools app, 1.1). Runs entirely in the browser on
                        pdf-lib + pdf.js, served from vendor/ and loaded the
                        first time the tab is opened — the other tabs pay
                        nothing for it. Results are written straight into the
                        shared export folder instead of Downloads.
Tab 4 — To Word:        The way back out (1.3), in-process and bundled like the
                        Documents tab. Born-digital PDFs are rebuilt as editable
                        .docx by pdf2docx — slow enough to need real progress, so
                        the conversion runs as a background job the page polls
                        page by page; a PDF with no selectable text is a scan, and
                        this tab says so rather than handing back a Word file full
                        of pictures of words. Markdown files are tokenised by
                        markdown-it and written out through python-docx with real
                        Word headings, lists, tables, quotes and links — including
                        the Obsidian dialect (frontmatter, wikilinks, embeds,
                        callouts, task boxes, ==highlights==).

Build: see build.bat (PyInstaller onefile; bundles markitdown, pdf-inspector,
pdf2docx, the OCR worker script and vendor/ as data). The OCR worker script lives
next to this file and is the single source of truth for it.
"""
import argparse
import atexit
import base64
import html
import json
import logging
import os
import re
import shutil
import subprocess
import sys
import tempfile
import threading
import time
import urllib.request
import uuid
import webbrowser
import zipfile
from collections import deque
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import urlparse, parse_qs

APP_NAME = "MarkdownFactory"
VERSION = "1.4"
PORT = 8770
WORKER_PORT = 8768  # keep OUT of the 8770-8779 double-launch scan range
DEFAULT_OUTPUT = Path.home() / "Desktop" / "MarkdownFactory-Export"

LOCAL_DIR = Path(os.environ.get("LOCALAPPDATA", Path.home())) / APP_NAME
RUNTIME_DIR = LOCAL_DIR / "ocr-runtime"
RUNTIME_PY = RUNTIME_DIR / "py" / "python.exe"
DEV_VENV_PY = Path.home() / ".venvs" / "unlimited-ocr" / "Scripts" / "python.exe"

EMBED_PY_URL = "https://www.python.org/ftp/python/3.12.8/python-3.12.8-embed-amd64.zip"
GETPIP_URL = "https://bootstrap.pypa.io/get-pip.py"
TORCH_ARGS = ["torch==2.8.0", "torchvision==0.23.0",
              "--index-url", "https://download.pytorch.org/whl/cu129"]
OCR_DEPS = ["transformers==4.57.1", "Pillow==12.1.1", "einops==0.8.2", "addict==2.4.0",
            "easydict==1.13", "pymupdf==1.27.2.2", "matplotlib==3.10.8", "psutil==7.2.2"]

WORKER_SOURCE_NAME = "unlimited-ocr-worker.py"
WORKER_EXTRA_ARGS = []  # dev: --ocr-no-model appends the worker's --no-model flag

# PDF tab: browser-side libraries, served on demand from vendor/
VENDOR_FILES = {
    "pdf-lib.min.js": "application/javascript",
    "pdf.min.js": "application/javascript",
    "pdf.worker.min.js": "application/javascript",
}


def asset_dir() -> Path:
    """Where bundled data lives: the PyInstaller temp dir, or this file's folder."""
    if getattr(sys, "frozen", False):
        return Path(sys._MEIPASS)  # noqa
    return Path(__file__).resolve().parent


# ---------------------------------------------------------------- documents tab

_md = None
_md_err = None
_md_lock = threading.Lock()

_pi = None
_pi_err = None


class UserFacingError(ValueError):
    """An error whose message is meant for the person who dropped the file, not
    for a log: a finished sentence that names the tab to try next. It subclasses
    ValueError because the To Word tab already treats a bare ValueError as
    'show this text as-is' (see run_word_job); the marker class lets the
    Documents endpoint make the same distinction explicitly rather than by
    catching a type half the standard library also raises."""


def get_markitdown():
    """Lazy import: the exe bundles markitdown, dev mode needs it pip-installed."""
    global _md, _md_err
    if _md is None and _md_err is None:
        try:
            from markitdown import MarkItDown
            _md = MarkItDown(enable_plugins=False)
        except Exception as e:
            _md_err = f"markitdown not available: {type(e).__name__}: {e}"
    if _md_err:
        raise RuntimeError(_md_err)
    return _md


def get_pdf_inspector():
    """Lazy import, same shape as get_markitdown(): pdf-inspector is a small Rust
    extension (one .pyd, no models, no network) that the exe bundles. If it isn't
    aboard — a dev machine without the wheel — PDFs simply take 1.3's road."""
    global _pi, _pi_err
    if _pi is None and _pi_err is None:
        try:
            import pdf_inspector
            _pi = pdf_inspector
        except Exception as e:
            _pi_err = f"pdf-inspector not available: {type(e).__name__}: {e}"
    if _pi_err:
        raise RuntimeError(_pi_err)
    return _pi


def unique_path(folder: Path, stem: str, ext: str = ".md") -> Path:
    p = folder / f"{stem}{ext}"
    n = 1
    while p.exists():
        n += 1
        p = folder / f"{stem} ({n}){ext}"
    return p


def skipped_pages_note(pages) -> str:
    """A mixed PDF's picture-only pages come back empty, and a reader who finds
    page 4 missing deserves to be told why rather than left to wonder.

    `pages` is PdfResult.pages_needing_ocr, which counts from 1. Note that
    classify_pdf()'s identically-named field counts from 0 — the two are not
    interchangeable, so this only ever reads the process_pdf() result."""
    if not pages:
        return ""
    if len(pages) == 1:
        which, noun, it, that = f"Page {pages[0]} of this PDF is", "a picture", "it", "that page"
    elif len(pages) <= 6:
        listed = ", ".join(str(p) for p in pages[:-1])
        which, noun, it, that = (f"Pages {listed} and {pages[-1]} of this PDF are",
                                 "pictures", "them", "those pages")
    else:
        which, noun, it, that = (f"{len(pages)} pages of this PDF are",
                                 "pictures", "them", "those pages")
    # "little or nothing" rather than "nothing": a picture page can still carry a
    # scrap of real text (bank statements put a printed mail-sort line above the
    # scanned image), and the note should not be caught overstating its case.
    return (f"> {which} {noun} rather than text, so little or nothing from {it} is "
            f"below. The Scans & Photos tab can read {that}.\n\n")


def convert_pdf(path: str, display_name: str = ""):
    """Markdown for a PDF via pdf-inspector, or None to let markitdown try.

    Raises UserFacingError when the file is a scan: markitdown hands back an
    empty .md for those, which looks like a conversion that worked until you
    open it. The To Word tab has turned scans away since 1.3; this is the
    Documents tab catching up.

    Anything else that goes wrong here — the engine missing, an encrypted or
    truncated file, a .pdf that isn't one — returns None and falls through to
    markitdown. 1.3's behaviour is the floor: 1.4 may be better than the old
    road but must never be worse than it."""
    try:
        result = get_pdf_inspector().process_pdf(path)
        # Read everything off the result inside the try as well: the fallback is
        # promised for "anything that goes wrong here", and an attribute that
        # moves in a future version of the engine should take the old road too,
        # not surface a raw AttributeError to someone who dropped a file.
        pdf_type = result.pdf_type
        text = result.markdown or ""
        ocr_pages = list(result.pages_needing_ocr or [])
    except Exception as e:
        # RuntimeError = no wheel aboard; ValueError = the engine looked and
        # refused (encrypted, corrupt, not a PDF). Neither is worth showing to
        # someone who just dropped a file — markitdown still gets its turn, and
        # gets to raise its own error if it fails too.
        # Name the file the user actually dropped, not the temp copy — a silent
        # fallback nobody can trace back to a document is not much of a log line.
        print(f"[documents] pdf-inspector passed on {display_name or Path(path).name} "
              f"({type(e).__name__}: {e}) — falling back to markitdown")
        return None

    # Raised outside the try above, so the fallback can never swallow the gate.
    if pdf_type in ("scanned", "image_based"):
        raise UserFacingError(
            "This PDF is a scan — its pages are pictures, with no text to read. "
            "Read it on the Scans & Photos tab instead.")

    if not text.strip():
        # Classified as readable but nothing came out. Rare, and markitdown may
        # still find something, so don't spend the user's file on our verdict.
        return None

    # Only a mixed document is part text and part picture, so only a mixed
    # document has missing pages to apologise for. On a text_based file the flag
    # means something closer to "this page has a figure in it", and a warning
    # about absent pages above a page whose text is right there would be a lie.
    note = skipped_pages_note(ocr_pages) if pdf_type == "mixed" else ""
    return note + text


def convert_document(filename: str, data: bytes, out_dir: Path) -> Path:
    suffix = Path(filename).suffix or ".bin"
    out_dir.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as tmp:
        tmp.write(data)
        tmp_path = tmp.name
    try:
        text = convert_pdf(tmp_path, filename) if suffix.lower() == ".pdf" else None
        if text is None:
            # Every other format, and any PDF the engine passed on. markitdown
            # is imported here rather than at the top of the function so a PDF
            # never fails for want of a library it didn't end up needing.
            md = get_markitdown()
            with _md_lock:
                result = md.convert(tmp_path)
            text = result.text_content or ""
        out_path = unique_path(out_dir, Path(filename).stem)
        out_path.write_text(text, encoding="utf-8")
        return out_path
    finally:
        try:
            os.unlink(tmp_path)
        except OSError:
            pass


# ---------------------------------------------------------------- PDF tab

def save_pdf(filename: str, data: bytes, out_dir: Path) -> Path:
    """The PDF tab does its work in the browser and posts the finished bytes here,
    so merged/split/compressed files land in the shared export folder rather than
    scattering through Downloads."""
    out_dir.mkdir(parents=True, exist_ok=True)
    stem = Path(filename).stem or "document"     # also strips any path trickery
    out_path = unique_path(out_dir, stem, ".pdf")
    out_path.write_bytes(data)
    return out_path


# ---------------------------------------------------------------- PDF -> Word tab

# Unlike the other conversions this one takes seconds per page, so it runs as a
# background job with a poll endpoint instead of one long-held POST: the browser
# gets page-by-page progress and a laptop that sleeps mid-convert doesn't lose a
# request it can never answer.
WORD_JOBS = {}                 # job id -> progress dict, newest last
WORD_JOBS_KEPT = 40
_word_lock = threading.Lock()  # one conversion at a time (see _ProgressHandler)
_word_jobs_lock = threading.Lock()
_p2d_err = None

_ANSI_RE = re.compile(r"\x1b\[[0-9;]*m")
_PHASE_RE = re.compile(r"^\[(\d)/4\]")
_PAGE_RE = re.compile(r"^\((\d+)/(\d+)\) Page\b")


def get_converter():
    """Lazy import: pdf2docx drags in PyMuPDF, OpenCV and NumPy — a second or two
    the first time. Only this tab pays it, and only once a file is dropped."""
    global _p2d_err
    if _p2d_err:
        raise RuntimeError(_p2d_err)
    try:
        from pdf2docx import Converter
        return Converter
    except Exception as e:
        _p2d_err = f"pdf2docx not available: {type(e).__name__}: {e}"
        raise RuntimeError(_p2d_err)


def _open_pdf(path: str):
    """PyMuPDF 1.24+ is importable under both names; pdf2docx still says fitz."""
    try:
        import pymupdf
    except ImportError:
        import fitz as pymupdf
    return pymupdf.open(path)


def probe_pdf(path: str):
    """(page_count, has_text). A PDF whose pages carry no extractable text is a
    scan — pdf2docx would return a Word file of page images, which is not what
    anyone means by 'convert to Word'."""
    doc = _open_pdf(path)
    try:
        chars = 0
        for i in range(min(10, doc.page_count)):
            chars += len(doc[i].get_text().strip())
            if chars >= 20:
                return doc.page_count, True
        return doc.page_count, chars >= 20
    finally:
        doc.close()


def parse_page_spec(spec: str, total: int):
    """'1-3, 5' -> [0, 1, 2, 4] (pdf2docx counts from zero). Blank -> None, the
    whole document. Raises ValueError with a sentence a non-technical user can
    act on, since this is free-text typed into a box."""
    spec = (spec or "").strip()
    if not spec:
        return None
    pages = []
    for chunk in spec.split(","):
        chunk = chunk.strip()
        if not chunk:
            continue
        m = re.fullmatch(r"(\d+)\s*(?:-\s*(\d+))?", chunk)
        if not m:
            raise ValueError(f"'{chunk}' isn't a page or a range — try 1-3, 5, 8-10.")
        lo = int(m.group(1))
        hi = int(m.group(2) or lo)
        if lo < 1 or hi > total or lo > hi:
            raise ValueError(f"'{chunk}' is outside this file's {total} pages.")
        pages.extend(range(lo - 1, hi))
    if not pages:
        return None
    return sorted(set(pages))


# ---- markdown -> Word ---------------------------------------------------
# The other direction from the Documents tab, and the one a colleague asks for
# when "can you send me that as a Word doc" means a note rather than a PDF.
# markdown-it-py tokenises; python-docx (already aboard for the PDF path) writes.

MD_SUFFIXES = {".md", ".markdown", ".mdown", ".mkd"}

_FRONTMATTER_RE = re.compile(r"\A﻿?---[ \t]*\r?\n.*?\r?\n---[ \t]*\r?\n", re.S)
_WIKILINK_RE = re.compile(r"(!?)\[\[([^\]|]+?)(?:\|([^\]]*))?\]\]")
_HIGHLIGHT_RE = re.compile(r"==(.+?)==", re.S)
_CALLOUT_RE = re.compile(r"^\[!([A-Za-z-]+)\][+-]?[ \t]*(.*)$")
_TASK_RE = re.compile(r"^\[([ xX])\][ \t]+")

CODE_FONT = "Consolas"
CODE_SHADE = "F4F4F4"
RULE_COLOR = "BFBFBF"
LINK_COLOR = (0x05, 0x63, 0xC1)
DIM_COLOR = (0x80, 0x80, 0x80)

_DOCX = {}


def _docx_bits():
    """Lazy, like every other engine here — nothing imports until a file lands."""
    if not _DOCX:
        from docx import Document
        from docx.enum.text import WD_COLOR_INDEX
        from docx.opc.constants import RELATIONSHIP_TYPE as RT
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Pt, RGBColor
        _DOCX.update(Document=Document, WD_COLOR_INDEX=WD_COLOR_INDEX, RT=RT,
                     OxmlElement=OxmlElement, qn=qn, Pt=Pt, RGBColor=RGBColor)
    return _DOCX


def decode_text(data: bytes) -> str:
    """Notes are UTF-8 in practice; anything older still has to open."""
    for encoding in ("utf-8-sig", "utf-8", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def obsidian_to_commonmark(text: str) -> str:
    """Obsidian's own syntax is not CommonMark, and three pieces of it matter:
    frontmatter is note metadata with nowhere to live in Word; a wikilink points
    into a vault Word cannot open, so it keeps only its visible label; an embed
    becomes an image node, which renders as one placeholder line because the
    browser uploaded the .md file and nothing around it."""
    text = text.replace("\r\n", "\n")
    text = _FRONTMATTER_RE.sub("", text, count=1)

    def link(m):
        bang, target, alias = m.group(1), m.group(2).strip(), m.group(3)
        label = (alias or "").strip() or target
        # embed targets are Obsidian filenames and nearly always contain spaces
        # ('Pasted image 20260730.png'), which a bare CommonMark destination
        # cannot hold — the angle-bracket form can
        return f"![{label}](<{target}>)" if bang else label

    return _WIKILINK_RE.sub(link, text)


class _MarkdownDocx:
    """markdown-it's token stream is flat, with open/close markers rather than a
    tree, so nesting rides on two small stacks instead of recursion."""

    def __init__(self, doc):
        self.doc = doc
        self.D = _docx_bits()
        self.lists = []          # 'ul' / 'ol', one per open list
        self.quotes = 0
        self.link_el = None      # runs go inside this while a link is open
        self.pending_task = False
        self.pending_callout = False

    # ---- inline runs

    def _style(self, run, marks):
        D = self.D
        if "strong" in marks:
            run.bold = True
        if "em" in marks:
            run.italic = True
        if "s" in marks:
            run.font.strike = True
        if "code" in marks:
            run.font.name = CODE_FONT
            run.font.size = D["Pt"](9.5)
        if "link" in marks:
            run.font.color.rgb = D["RGBColor"](*LINK_COLOR)
            run.underline = True
        if "dim" in marks:
            run.italic = True
            run.font.color.rgb = D["RGBColor"](*DIM_COLOR)

    def _run(self, para, text, marks):
        run = para.add_run(text)
        self._style(run, marks)
        if self.link_el is not None:
            self.link_el.append(run._r)   # a run only counts as linked in here
        return run

    def _text(self, para, text, marks):
        """==highlight== is Obsidian's, not CommonMark's, so it is applied at the
        run level — which makes it work inside bold, links and table cells too."""
        pos = 0
        for m in _HIGHLIGHT_RE.finditer(text):
            if m.start() > pos:
                self._run(para, text[pos:m.start()], marks)
            run = self._run(para, m.group(1), marks)
            run.font.highlight_color = self.D["WD_COLOR_INDEX"].YELLOW
            pos = m.end()
        if pos < len(text):
            self._run(para, text[pos:], marks)

    def _open_link(self, para, href):
        D = self.D
        if not href or href.startswith("#"):
            return None
        try:
            r_id = para.part.relate_to(href, D["RT"].HYPERLINK, is_external=True)
        except Exception:
            return None      # a malformed href is not worth losing the text over
        el = D["OxmlElement"]("w:hyperlink")
        el.set(D["qn"]("r:id"), r_id)
        para._p.append(el)
        return el

    def inline(self, para, children):
        marks = []
        for t in children or []:
            ty = t.type
            if ty == "text":
                self._text(para, t.content, marks)
            elif ty == "code_inline":
                self._text(para, t.content, marks + ["code"])
            elif ty in ("strong_open", "em_open", "s_open"):
                marks.append(ty[:-5])
            elif ty in ("strong_close", "em_close", "s_close"):
                name = ty[:-6]
                if name in marks:
                    marks.remove(name)
            elif ty == "link_open":
                self.link_el = self._open_link(para, t.attrGet("href") or "")
                marks.append("link")
            elif ty == "link_close":
                self.link_el = None
                if "link" in marks:
                    marks.remove("link")
            elif ty == "softbreak":
                self._run(para, " ", marks)
            elif ty == "hardbreak":
                self._run(para, "", marks).add_break()
            elif ty == "image":
                label = (t.content or "").strip() or (t.attrGet("src") or "image")
                self._run(para, f"[image: {label}]", marks + ["dim"])
            elif ty == "html_inline":
                self._text(para, t.content, marks)

    # ---- block shells

    def _paragraph(self):
        if self.lists:
            depth = min(len(self.lists), 3)
            base = "List Number" if self.lists[-1] == "ol" else "List Bullet"
            return self.doc.add_paragraph(style=base if depth == 1 else f"{base} {depth}")
        if self.quotes:
            return self.doc.add_paragraph(style="Quote")
        return self.doc.add_paragraph()

    def _shade(self, para):
        D = self.D
        shd = D["OxmlElement"]("w:shd")
        shd.set(D["qn"]("w:val"), "clear")
        shd.set(D["qn"]("w:fill"), CODE_SHADE)
        para._p.get_or_add_pPr().append(shd)

    def _rule(self):
        D = self.D
        para = self.doc.add_paragraph()
        bdr = D["OxmlElement"]("w:pBdr")
        bottom = D["OxmlElement"]("w:bottom")
        for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"),
                     ("w:color", RULE_COLOR)):
            bottom.set(D["qn"](k), v)
        bdr.append(bottom)
        para._p.get_or_add_pPr().append(bdr)

    def _code_block(self, content):
        D = self.D
        for line in (content or "").rstrip("\n").split("\n"):
            para = self.doc.add_paragraph()
            para.paragraph_format.space_before = D["Pt"](0)
            para.paragraph_format.space_after = D["Pt"](0)
            para.paragraph_format.left_indent = D["Pt"](18)
            run = para.add_run(line)
            run.font.name = CODE_FONT
            run.font.size = D["Pt"](9)
            self._shade(para)
        self.doc.add_paragraph().paragraph_format.space_after = D["Pt"](0)

    def _checkbox(self, para, children):
        """'- [ ] thing' is a plain list item to CommonMark; Word has no live
        checkbox either, so it gets the ballot-box glyphs Word renders."""
        if not (children and children[0].type == "text"):
            return
        m = _TASK_RE.match(children[0].content)
        if not m:
            return
        self._run(para, "☑ " if m.group(1).lower() == "x" else "☐ ", [])
        children[0].content = children[0].content[m.end():]

    def _callout(self, para, children):
        """Obsidian callouts ('> [!warning] Contradiction') are blockquotes with a
        typed first line. Word gets the type and title as a bold lead-in."""
        if not (children and children[0].type == "text"):
            return
        m = _CALLOUT_RE.match(children[0].content)
        if not m:
            return
        kind, title = m.group(1).upper(), m.group(2).strip()
        self._run(para, f"{kind}{' — ' + title if title else ''}", ["strong"])
        children.pop(0)
        if children and children[0].type == "softbreak":
            children.pop(0)
        if children:
            para.add_run().add_break()

    def _table(self, tokens, i):
        """Consume table_open..table_close and return the index after it."""
        rows, j = [], i + 1
        while j < len(tokens) and tokens[j].type != "table_close":
            if tokens[j].type == "tr_open":
                cells = []
                j += 1
                while j < len(tokens) and tokens[j].type != "tr_close":
                    if tokens[j].type in ("th_open", "td_open"):
                        head = tokens[j].type == "th_open"
                        nxt = tokens[j + 1] if j + 1 < len(tokens) else None
                        cells.append((head, nxt if nxt is not None and
                                      nxt.type == "inline" else None))
                    j += 1
                rows.append(cells)
            j += 1

        if rows:
            table = self.doc.add_table(rows=0, cols=max(len(r) for r in rows))
            table.style = "Table Grid"
            for cells in rows:
                row = table.add_row()
                for k, (head, tok) in enumerate(cells):
                    if k >= len(row.cells):
                        break
                    para = row.cells[k].paragraphs[0]
                    self.link_el = None      # links never span cells
                    self.inline(para, tok.children if tok is not None else [])
                    self.link_el = None
                    if head:
                        for run in para.runs:
                            run.bold = True
            self.doc.add_paragraph()
        return j + 1

    # ---- the walk

    def render(self, tokens):
        i = 0
        while i < len(tokens):
            t = tokens[i]
            ty = t.type

            if ty == "heading_open":
                para = self.doc.add_paragraph(style=f"Heading {min(int(t.tag[1:]), 6)}")
                if i + 1 < len(tokens) and tokens[i + 1].type == "inline":
                    self.inline(para, tokens[i + 1].children)
                i += 3
                continue

            if ty == "paragraph_open":
                para = self._paragraph()
                children = (tokens[i + 1].children
                            if i + 1 < len(tokens) and tokens[i + 1].type == "inline"
                            else [])
                if self.pending_task:
                    self.pending_task = False
                    self._checkbox(para, children)
                if self.pending_callout:
                    self.pending_callout = False
                    self._callout(para, children)
                self.inline(para, children)
                i += 3
                continue

            if ty in ("bullet_list_open", "ordered_list_open"):
                self.lists.append("ol" if ty.startswith("ordered") else "ul")
            elif ty in ("bullet_list_close", "ordered_list_close"):
                if self.lists:
                    self.lists.pop()
            elif ty == "list_item_open":
                self.pending_task = True
            elif ty == "blockquote_open":
                self.quotes += 1
                self.pending_callout = True
            elif ty == "blockquote_close":
                self.quotes = max(0, self.quotes - 1)
                self.pending_callout = False
            elif ty in ("fence", "code_block"):
                self._code_block(t.content)
            elif ty == "hr":
                self._rule()
            elif ty == "html_block":
                for line in (t.content or "").strip().split("\n"):
                    if line.strip():
                        self._run(self.doc.add_paragraph(), line, ["dim"])
            elif ty == "table_open":
                i = self._table(tokens, i)
                continue
            i += 1


def convert_markdown(text: str, out_path: Path):
    from markdown_it import MarkdownIt
    body = obsidian_to_commonmark(text)
    if not body.strip():
        raise ValueError("There is nothing in this file to put into Word.")
    # commonmark, not gfm-like: gfm-like turns on linkify, which needs a package
    # that is not aboard. Tables and ~~strikethrough~~ are the parts worth having.
    md = MarkdownIt("commonmark").enable(["table", "strikethrough"])
    doc = _docx_bits()["Document"]()
    _MarkdownDocx(doc).render(md.parse(body))
    doc.save(str(out_path))


class _ProgressHandler(logging.Handler):
    """pdf2docx narrates itself onto the root logger ('[3/4] Parsing pages...',
    '(2/16) Page 2') and offers no callback, so reading the log back is the only
    page-level progress there is. Safe because conversions hold _word_lock: one
    handler is attached at a time, and it only ever writes to its own job."""

    def __init__(self, job: dict):
        super().__init__(level=logging.INFO)
        self.job = job

    def emit(self, record):
        try:
            msg = _ANSI_RE.sub("", record.getMessage()).strip()
        except Exception:
            return
        m = _PHASE_RE.match(msg)
        if m:
            self.job["phase"] = int(m.group(1))
            return
        m = _PAGE_RE.match(msg)
        if m:
            self.job["page"] = int(m.group(1))
            self.job["pages"] = int(m.group(2))


def _word_note(job: dict) -> str:
    """Plain-English status line, derived from whatever the log last said."""
    phase, page, pages = job.get("phase", 0), job.get("page", 0), job.get("pages", 0)
    if job["state"] == "queued":
        return "waiting for the file ahead of it…"
    if job.get("kind") == "md":
        return "writing the Word file…"   # markdown is over in milliseconds
    if phase >= 4 and pages:
        return f"building the Word file — page {page} of {pages}"
    if phase == 3 and pages:
        return f"reading the layout — page {page} of {pages}"
    if phase == 2:
        return "analysing the document…"
    return "opening the file…"


def _word_pct(job: dict) -> int:
    """Reading and writing are two passes over every page, so each is half the bar."""
    phase, page, pages = job.get("phase", 0), job.get("page", 0), job.get("pages", 0)
    if job["state"] == "done":
        return 100
    if not pages:
        return 3
    if phase >= 4:
        return min(99, 50 + round(50 * page / pages))
    if phase == 3:
        return min(49, round(47 * page / pages))
    return 3


def new_word_job(filename: str) -> str:
    job_id = uuid.uuid4().hex[:12]
    kind = "md" if Path(filename).suffix.lower() in MD_SUFFIXES else "pdf"
    with _word_jobs_lock:
        WORD_JOBS[job_id] = {"state": "queued", "name": filename, "kind": kind,
                             "phase": 0, "page": 0, "pages": 0, "path": "", "error": ""}
        for stale in list(WORD_JOBS)[:-WORD_JOBS_KEPT]:
            WORD_JOBS.pop(stale, None)
    return job_id


def run_word_job(job_id: str, filename: str, data: bytes, out_dir: Path, page_spec: str):
    job = WORD_JOBS[job_id]
    tmp_path = None
    out_path = None
    suffix = Path(filename).suffix.lower()
    try:
        if suffix in MD_SUFFIXES:
            # Markdown is parsed in memory and written straight out — no engine to
            # warm up, no temp file, and nothing to report progress on. It still
            # takes _word_lock so that two files can never race for the same name.
            text = decode_text(data)
            with _word_lock:
                job["state"] = "running"
                out_dir.mkdir(parents=True, exist_ok=True)
                out_path = unique_path(out_dir, Path(filename).stem, ".docx")
                convert_markdown(text, out_path)
        elif suffix == ".pdf":
            Converter = get_converter()
            with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as tmp:
                tmp.write(data)
                tmp_path = tmp.name

            total, has_text = probe_pdf(tmp_path)
            if not has_text:
                raise ValueError("This PDF is a scan — its pages are pictures, with no "
                                 "text to put into Word. Read it on the Scans & Photos "
                                 "tab first.")
            pages = parse_page_spec(page_spec, total)

            with _word_lock:
                job["state"] = "running"
                out_dir.mkdir(parents=True, exist_ok=True)
                out_path = unique_path(out_dir, Path(filename).stem, ".docx")
                cv = Converter(tmp_path)
                handler = _ProgressHandler(job)
                root = logging.getLogger()
                was = root.level            # pdf2docx sets this to INFO on import, but
                root.setLevel(logging.INFO)  # don't let a stray basicConfig mute progress
                root.addHandler(handler)
                try:
                    cv.convert(str(out_path), pages=pages)
                finally:
                    root.removeHandler(handler)
                    root.setLevel(was)
                    cv.close()
        else:
            raise ValueError("This tab turns PDFs and Markdown files into Word — use "
                             "the Documents tab for Word, PowerPoint, Excel and web pages.")
        job.update(state="done", path=str(out_path))
    except Exception as e:
        if out_path is not None and out_path.exists():
            try:
                out_path.unlink()   # never leave a half-written .docx behind
            except OSError:
                pass
        msg = str(e) if isinstance(e, ValueError) else f"{type(e).__name__}: {e}"
        job.update(state="error", error=msg)
    finally:
        if tmp_path:
            try:
                os.unlink(tmp_path)
            except OSError:
                pass


# ---------------------------------------------------------------- OCR tab manager

OCR = {"state": "checking", "detail": "", "log": deque(maxlen=300)}
_ocr_lock = threading.Lock()
_worker_proc = None


def _log(line: str):
    line = line.rstrip()
    if line:
        OCR["log"].append(line)
        OCR["detail"] = line


def worker_source() -> Path:
    """The copy that ships with this project (bundled as data in the exe). Since
    1.1 this is the single source of truth — the standalone Unlimited-OCR GUI it
    came from is shelved, so there is no sibling project to read from."""
    return asset_dir() / WORKER_SOURCE_NAME


def has_nvidia_gpu() -> bool:
    exe = shutil.which("nvidia-smi")
    if not exe:
        return False
    try:
        return subprocess.run([exe], capture_output=True, timeout=15).returncode == 0
    except Exception:
        return False


def probe_python(py: Path) -> bool:
    """True if this interpreter has CUDA torch (first torch import can be slow)."""
    if not py.exists():
        return False
    try:
        r = subprocess.run(
            [str(py), "-c", "import torch,sys;sys.exit(0 if torch.cuda.is_available() else 3)"],
            capture_output=True, timeout=180,
        )
        return r.returncode == 0
    except Exception:
        return False


def find_runtime_python():
    for py in (DEV_VENV_PY, RUNTIME_PY):
        if probe_python(py):
            return py
    return None


def worker_alive() -> bool:
    try:
        with urllib.request.urlopen(f"http://127.0.0.1:{WORKER_PORT}/status", timeout=3) as r:
            json.loads(r.read())
        return True
    except Exception:
        return False


def start_worker(py: Path):
    """Run the Unlimited-OCR GUI server as a child; its own page handles the rest
    (model download banner, progress pills, cancel, checkpoints)."""
    global _worker_proc
    if worker_alive():
        OCR.update(state="ready")
        _log("OCR worker already running — reusing it")
        return
    src = worker_source()
    if not src.exists():
        OCR.update(state="error")
        _log(f"worker source missing: {src}")
        return
    run_copy = RUNTIME_DIR / WORKER_SOURCE_NAME
    RUNTIME_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copy2(src, run_copy)
    OCR.update(state="starting")
    _log("starting OCR engine…")
    flags = subprocess.CREATE_NO_WINDOW if sys.platform == "win32" else 0
    _worker_proc = subprocess.Popen(
        [str(py), str(run_copy), "--no-browser", "--port", str(WORKER_PORT),
         "--output-dir", str(DEFAULT_OUTPUT)]   # both tabs share one export folder
        + WORKER_EXTRA_ARGS,
        creationflags=flags,
    )
    for _ in range(60):
        if worker_alive():
            OCR.update(state="ready")
            _log("OCR engine is up — model loads in its pane")
            return
        if _worker_proc.poll() is not None:
            break
        time.sleep(1)
    OCR.update(state="error")
    _log("OCR engine failed to start")


def _stop_worker():
    if _worker_proc and _worker_proc.poll() is None:
        _worker_proc.terminate()


atexit.register(_stop_worker)


def _download(url: str, dest: Path, label: str):
    dest.parent.mkdir(parents=True, exist_ok=True)
    _log(f"downloading {label}…")

    def hook(blocks, bs, total):
        if total > 0:
            pct = min(100, blocks * bs * 100 // total)
            OCR["detail"] = f"downloading {label}… {pct}%"

    urllib.request.urlretrieve(url, dest, reporthook=hook)
    _log(f"downloaded {label}")


def _run_logged(args, label: str):
    _log(f"{label}…")
    p = subprocess.Popen(args, stdout=subprocess.PIPE, stderr=subprocess.STDOUT,
                         text=True, encoding="utf-8", errors="replace")
    for line in p.stdout:
        _log(line)
    if p.wait() != 0:
        raise RuntimeError(f"{label} failed — see log")


def setup_ocr_runtime():
    """One-time bootstrap: embedded Python + CUDA torch + OCR deps into LOCALAPPDATA.
    ~4 GB download; every step streams to the log the page polls."""
    try:
        OCR.update(state="installing")
        py_dir = RUNTIME_DIR / "py"
        if not RUNTIME_PY.exists():
            zip_path = RUNTIME_DIR / "python-embed.zip"
            _download(EMBED_PY_URL, zip_path, "Python runtime (11 MB)")
            py_dir.mkdir(parents=True, exist_ok=True)
            with zipfile.ZipFile(zip_path) as z:
                z.extractall(py_dir)
            zip_path.unlink()
            for pth in py_dir.glob("python*._pth"):  # enable site-packages
                pth.write_text(pth.read_text().replace("#import site", "import site"),
                               encoding="ascii")
        getpip = RUNTIME_DIR / "get-pip.py"
        _download(GETPIP_URL, getpip, "package installer")
        _run_logged([str(RUNTIME_PY), str(getpip), "--no-warn-script-location"],
                    "installing pip")
        _run_logged([str(RUNTIME_PY), "-m", "pip", "install", "--no-warn-script-location"]
                    + TORCH_ARGS, "installing GPU libraries (~4 GB, the big one)")
        _run_logged([str(RUNTIME_PY), "-m", "pip", "install", "--no-warn-script-location"]
                    + OCR_DEPS, "installing OCR libraries")
        _log("verifying GPU access…")
        if not probe_python(RUNTIME_PY):
            raise RuntimeError("installed, but torch cannot see a CUDA GPU")
        start_worker(RUNTIME_PY)
    except Exception as e:
        OCR.update(state="error")
        _log(f"setup failed: {type(e).__name__}: {e}")


def init_ocr():
    """Startup: figure out what state the OCR tab should show."""
    with _ocr_lock:
        if worker_alive():
            OCR.update(state="ready")
            _log("OCR worker already running — reusing it")
            return
        if not has_nvidia_gpu():
            OCR.update(state="no-gpu")
            _log("no NVIDIA GPU found")
            return
        py = find_runtime_python()
        if py:
            start_worker(py)
        else:
            OCR.update(state="needs-setup")
            _log("OCR runtime not installed yet")


# ---------------------------------------------------------------- page

HTML_PAGE = """<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>MarkdownFactory v__VERSION__</title>
<!-- 1.2 favicon: the logo reduced to the three marks that survive 16px —
     black field, gold frame (#ffde59), red wordmark bar (#ff3131), sky rule
     (#38b6ff). Inline data URI, so no second file and no /favicon.ico route;
     same composition and stroke weights MangaFactory settled on, in this
     app's own inks. The logo's own hairlines (frame ~1.5, rule ~0.7 of 32)
     rasterise to nothing at 16px — 3 keeps every mark its own colour at the
     pixel centre. -->
<link rel="icon" href="data:image/svg+xml,%3Csvg xmlns='http://www.w3.org/2000/svg' viewBox='0 0 32 32'%3E%3Crect width='32' height='32' fill='%23000000'/%3E%3Crect x='4.5' y='4.5' width='23' height='23' fill='none' stroke='%23ffde59' stroke-width='3'/%3E%3Crect x='9' y='9.5' width='14' height='4.5' fill='%23ff3131'/%3E%3Crect x='9' y='17' width='3' height='7.5' fill='%2338b6ff'/%3E%3C/svg%3E">
<link rel="preconnect" href="https://fonts.googleapis.com">
<link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
<!-- Loaded non-render-blocking: this app is offline-first, so a missing
     network must never delay first paint. Offline you get the fallbacks. -->
<link rel="stylesheet" media="print" onload="this.media='all'"
      href="https://fonts.googleapis.com/css2?family=Space+Grotesk:wght@400;500;600;700&family=JetBrains+Mono:wght@400;500;700&display=swap">
<style>
  :root {
    /* Palette lifted from MarkDownFactory Logo 1.0 — black field, red
       wordmark, gold frame, thin blue rule. Flat and dark, MangaFactory's
       rules inverted: no gradients, no glows, no coloured shadows anywhere. */
    --paper: #000000;        /* logo background — page bg */
    --paper-2: #0b0b0d;      /* wells, log boxes, thumbnail beds */
    --card: #121215;         /* raised surface */
    --field: #1a1a1e;        /* input + dropzone backgrounds */
    --ink: #f2f0ee;          /* primary text — 18.9:1 on paper */
    --ink-soft: #b4b0ac;     /* secondary text — 9.7:1 */
    --muted: #8a8681;        /* meta text — 5.8:1 on paper, 5.2:1 on card */
    --line: #26262b;         /* hairline borders */
    --line-strong: #3c3c44;  /* emphasized borders */

    /* The logo's marks. Each tab owns one: it re-points --accent below, so
       every control inside a pane picks up that tab's colour for free. The
       logo carries three inks and a page; the fourth tab takes the page. */
    --red:  #ff3131;         /* wordmark — Documents, and the app's identity */
    --gold: #ffde59;         /* frame — Scans & Photos */
    --sky:  #38b6ff;         /* rule — PDF Tools */
    --page: #f2f0ee;         /* the paper itself — PDF to Word */
    --accent: var(--red);
    --accent-soft: rgba(255,49,49,0.10);
    /* Every logo ink is bright (5.7:1, 15.8:1, 9.3:1, 18.9:1 on black), so
       filled controls take dark text — white would fail AA on every one. */
    --on-accent: #0d0a0a;

    --success: #62c98c; --warn: var(--gold); --danger: #ff8080;
    --mono: 'JetBrains Mono', 'Cascadia Mono', Consolas, monospace;
    --sans: 'Space Grotesk', 'Segoe UI', system-ui, sans-serif;
    --r: 2px;
  }
  * { box-sizing: border-box; margin: 0; padding: 0; }
  body {
    background: var(--paper); color: var(--ink); font-family: var(--sans);
    min-height: 100vh; overflow-x: hidden; -webkit-font-smoothing: antialiased;
  }
  .container { max-width: 900px; margin: 0 auto; padding: 44px 24px 88px; }
  code { font-family: var(--mono); font-size: .92em; color: var(--ink-soft); }

  /* Header — a miniature of the logo: gold frame, red wordmark, blue rule */
  header {
    display: flex; align-items: flex-end; gap: 18px; flex-wrap: wrap;
    margin-bottom: 28px; padding-bottom: 22px; border-bottom: 1px solid var(--line);
  }
  .logo-frame { border: 1.5px solid var(--gold); border-radius: var(--r); padding: 12px 18px 10px; flex-shrink: 0; }
  .logo-word { font-size: 24px; font-weight: 700; letter-spacing: 1px; color: var(--red); line-height: 1; }
  .logo-sub { display: flex; align-items: center; gap: 7px; margin-top: 5px; font-size: 10px; font-weight: 600; letter-spacing: 4px; color: var(--red); }
  .logo-rule { width: 1.5px; height: 12px; background: var(--sky); flex-shrink: 0; }
  .tagline { font-size: 12px; color: var(--muted); letter-spacing: .4px; padding-bottom: 3px; }
  .version { font-family: var(--mono); font-size: 11px; color: var(--red); margin-left: auto; letter-spacing: 2px; padding-bottom: 3px; }

  /* Tabs — flat underline, one logo ink each */
  .tabs { display: flex; gap: 26px; flex-wrap: wrap; margin-bottom: 20px; border-bottom: 1px solid var(--line); }
  .tab {
    background: none; border: none; border-bottom: 2px solid transparent; margin-bottom: -1px;
    padding: 10px 2px 12px; cursor: pointer; font-family: var(--sans); font-size: 13px;
    font-weight: 600; letter-spacing: 1.5px; text-transform: uppercase; color: var(--muted);
    transition: color .15s, border-color .15s;
  }
  .tab:hover { color: var(--ink); }
  .tab[data-tab="docs"].active { color: var(--red);  border-bottom-color: var(--red); }
  .tab[data-tab="ocr"].active  { color: var(--gold); border-bottom-color: var(--gold); }
  .tab[data-tab="pdf"].active  { color: var(--sky);  border-bottom-color: var(--sky); }
  .tab[data-tab="word"].active { color: var(--page); border-bottom-color: var(--page); }
  .tab-content { display: none; }
  .tab-content.active { display: block; animation: fadeIn .25s ease-out; }
  @keyframes fadeIn { from { opacity: 0; transform: translateY(6px); } to { opacity: 1; transform: none; } }

  #tab-docs { --accent: var(--red);  --accent-soft: rgba(255,49,49,.10); }
  #tab-ocr  { --accent: var(--gold); --accent-soft: rgba(255,222,89,.10); }
  #tab-pdf  { --accent: var(--sky);  --accent-soft: rgba(56,182,255,.10); }
  #tab-word { --accent: var(--page); --accent-soft: rgba(242,240,238,.08); }

  /* Shared surfaces */
  .card { background: var(--card); border: 1px solid var(--line); border-radius: var(--r); padding: 22px; }
  .card-title { font-size: 11px; font-weight: 600; letter-spacing: 2.2px; text-transform: uppercase; color: var(--ink); padding-bottom: 10px; border-bottom: 1px solid var(--line); margin-bottom: 16px; }
  .sub { font-size: 13px; color: var(--muted); line-height: 1.7; margin-bottom: 16px; max-width: 70ch; }
  .hidden { display: none; }

  /* Output folder — shared by every tab but Scans & Photos */
  .outrow { display: flex; gap: 10px; align-items: center; margin-bottom: 18px; }
  .outrow label { font-family: var(--mono); font-size: 11px; letter-spacing: .5px; color: var(--muted); white-space: nowrap; }
  .outrow input {
    flex: 1; min-width: 0; background: var(--field); border: 1px solid var(--line-strong);
    border-radius: var(--r); color: var(--ink); font-family: var(--mono); font-size: 12px;
    padding: 9px 12px; outline: none; transition: border-color .15s;
  }
  .outrow input:focus { border-color: var(--accent); }

  /* Buttons */
  .go {
    background: var(--accent); color: var(--on-accent); border: none; border-radius: var(--r);
    font-family: var(--sans); font-size: 14px; font-weight: 600; padding: 12px 22px;
    cursor: pointer; white-space: nowrap; transition: filter .15s, background .15s, color .15s;
  }
  .go:not(:disabled):hover { filter: brightness(1.12); }
  .go:not(:disabled):active { filter: brightness(.92); }
  .go:disabled { background: var(--field); color: var(--muted); cursor: not-allowed; }
  .ghost {
    background: transparent; border: 1px solid var(--line-strong); color: var(--ink);
    border-radius: var(--r); font-family: var(--sans); font-size: 13px; padding: 11px 16px;
    cursor: pointer; white-space: nowrap; transition: border-color .15s, color .15s, background .15s;
  }
  .ghost:not(:disabled):hover { border-color: var(--accent); color: var(--accent); background: var(--accent-soft); }
  .ghost:disabled { opacity: .4; cursor: not-allowed; }
  .btn-sm { font-size: 12px; padding: 8px 14px; }

  /* Dropzones — one look for all five */
  #drop, .dropzone {
    border: 1px dashed var(--line-strong); border-radius: var(--r); background: var(--field);
    padding: 34px 20px; text-align: center; cursor: pointer;
    transition: border-color .15s, background .15s;
  }
  #drop:hover, .dropzone:hover { border-color: var(--accent); }
  #drop.over, .dropzone.drag { border-color: var(--accent); background: var(--accent-soft); }
  #drop .big, .dropzone .big { font-size: 14px; font-weight: 600; letter-spacing: .3px; color: var(--ink); }
  #drop .small, .dropzone .small { font-family: var(--mono); font-size: 11px; color: var(--muted); margin-top: 7px; line-height: 1.6; }
  #fileInput { display: none; }

  /* Documents tab — conversion queue. Since 1.4 the row grows a line of text
     under the name, the same .note the To Word rows use: the scan gate has a
     sentence to say, and a sentence that lives only in a tooltip may as well
     not exist. A row with nothing to add stays exactly as tall as it was. */
  #tab-docs ul { list-style: none; margin-top: 16px; display: flex; flex-direction: column; gap: 6px; }
  #tab-docs li {
    display: flex; flex-direction: column; gap: 7px; padding: 11px 14px;
    background: var(--field); border: 1px solid var(--line); border-radius: var(--r);
  }
  #tab-docs li .top { display: flex; align-items: center; gap: 12px; }
  #tab-docs li .note:empty { display: none; }
  #tab-docs li .name {
    flex: 1; min-width: 0; font-family: var(--mono); font-size: 12px; color: var(--ink);
    overflow: hidden; text-overflow: ellipsis; white-space: nowrap;
  }
  .st {
    font-family: var(--mono); font-size: 9px; font-weight: 700; letter-spacing: 1px;
    text-transform: uppercase; padding: 3px 9px; border: 1px solid; border-radius: var(--r);
    white-space: nowrap;
  }
  .st.wait { color: var(--muted);   border-color: var(--line-strong); }
  .st.run  { color: var(--sky);     border-color: rgba(56,182,255,.45);  background: rgba(56,182,255,.10); }
  .st.ok   { color: var(--success); border-color: rgba(98,201,140,.45);  background: rgba(98,201,140,.12); }
  .st.err  { color: var(--danger);  border-color: rgba(255,128,128,.45); background: rgba(255,128,128,.12); cursor: help; }
  #summary { font-family: var(--mono); font-size: 11px; color: var(--muted); margin-top: 14px; }

  /* PDF to Word tab — same queue, but each row grows a progress bar and a
     status line, because a long PDF is minutes of work rather than seconds. */
  #tab-word ul { list-style: none; margin-top: 16px; display: flex; flex-direction: column; gap: 6px; }
  #tab-word li {
    display: flex; flex-direction: column; gap: 9px; padding: 12px 14px;
    background: var(--field); border: 1px solid var(--line); border-radius: var(--r);
  }
  #tab-word li .top { display: flex; align-items: center; gap: 12px; }
  #tab-word li .name {
    flex: 1; min-width: 0; font-family: var(--mono); font-size: 12px; color: var(--ink);
    overflow: hidden; text-overflow: ellipsis; white-space: nowrap;
  }
  .bar { height: 3px; background: var(--paper-2); border-radius: var(--r); overflow: hidden; }
  .bar > i { display: block; height: 100%; width: 0; background: var(--accent); transition: width .3s ease-out; }
  .bar.done > i { background: var(--success); }
  .bar.failed > i { background: var(--danger); }
  .note {
    font-family: var(--mono); font-size: 11px; color: var(--muted);
    line-height: 1.6; overflow-wrap: anywhere;
  }
  .note.failed { color: var(--danger); }
  /* The parenthetical does the distinguishing, not a dimmer grey — anything
     below --muted would break the theme's 4.8:1 floor. */
  .hint { color: var(--muted); text-transform: none; letter-spacing: 0; }
  #wordSummary { font-family: var(--mono); font-size: 11px; color: var(--muted); margin-top: 14px; }

  /* Scans & Photos tab */
  #ocrTitle { font-size: 15px; font-weight: 700; color: var(--ink); margin-bottom: 10px; }
  #ocrCard p { font-size: 13px; color: var(--ink-soft); line-height: 1.7; max-width: 66ch; }
  #setupBtn { margin-top: 16px; }
  #setupStage { font-family: var(--mono); font-size: 12px; color: var(--accent); margin-top: 14px; min-height: 1.2em; }
  #setupLog {
    display: none; background: var(--paper-2); border: 1px solid var(--line); border-radius: var(--r);
    margin-top: 12px; padding: 14px 16px; font-family: var(--mono); font-size: 11px;
    color: var(--ink-soft); line-height: 1.8; max-height: 15rem; overflow-y: auto; white-space: pre-wrap;
  }
  #setupLog::-webkit-scrollbar { width: 6px; }
  #setupLog::-webkit-scrollbar-track { background: transparent; }
  #setupLog::-webkit-scrollbar-thumb { background: var(--line-strong); border-radius: var(--r); }
  /* The worker serves its own page, so this pane is the one seam in the
     theme — the hairline makes it read as a deliberate frame. */
  #ocrFrame { width: 100%; height: 78vh; border: 1px solid var(--line); border-radius: var(--r); background: var(--paper); display: none; }

  /* PDF Tools tab */
  .subtabs { display: flex; gap: 20px; flex-wrap: wrap; margin-bottom: 18px; border-bottom: 1px solid var(--line); }
  .subtab {
    background: none; border: none; border-bottom: 2px solid transparent; margin-bottom: -1px;
    padding: 8px 2px 10px; cursor: pointer; font-family: var(--sans); font-size: 12px;
    font-weight: 600; letter-spacing: 1.2px; text-transform: uppercase; color: var(--muted);
    transition: color .15s, border-color .15s;
  }
  .subtab:hover { color: var(--ink); }
  .subtab.active { color: var(--accent); border-bottom-color: var(--accent); }
  #pdfBoot { font-family: var(--mono); font-size: 11px; color: var(--muted); margin-bottom: 14px; }
  #pdfBoot:empty { display: none; }
  #pdfBoot.err { color: var(--danger); }

  .list { list-style: none; margin-top: 16px; display: flex; flex-direction: column; gap: 6px; }
  .item {
    display: flex; align-items: center; gap: 12px; background: var(--field);
    border: 1px solid var(--line); border-radius: var(--r); padding: 10px 12px;
    cursor: grab; transition: border-color .15s, opacity .15s;
  }
  .item.dragging { opacity: .35; }
  .item.over { border-color: var(--accent); }
  .grip { color: var(--muted); font-size: 15px; line-height: 1; user-select: none; }
  .idx {
    flex: 0 0 24px; height: 24px; border-radius: var(--r); background: var(--paper-2);
    border: 1px solid var(--line); color: var(--muted); font-family: var(--mono);
    font-size: 11px; font-weight: 700; display: flex; align-items: center; justify-content: center;
  }
  .thumb {
    flex: 0 0 auto; width: 42px; height: 56px; border-radius: var(--r);
    background: var(--paper-2); border: 1px solid var(--line); cursor: zoom-in;
    display: flex; align-items: center; justify-content: center; overflow: hidden;
    color: var(--muted); font-size: 15px; transition: border-color .15s;
  }
  .thumb:hover { border-color: var(--accent); }
  .thumb img { width: 100%; height: 100%; object-fit: cover; object-position: top center; }
  .thumb.loading { font-size: 11px; }
  .meta { flex: 1 1 auto; min-width: 0; }
  .meta .name { font-family: var(--mono); font-size: 12px; color: var(--ink); white-space: nowrap; overflow: hidden; text-overflow: ellipsis; }
  .info { font-family: var(--mono); font-size: 10px; color: var(--muted); margin-top: 3px; }
  .btns { display: flex; gap: 4px; }
  .iconbtn {
    background: transparent; border: 1px solid var(--line-strong); color: var(--ink-soft);
    width: 28px; height: 28px; border-radius: var(--r); cursor: pointer; font-size: 12px;
    display: flex; align-items: center; justify-content: center;
    transition: border-color .15s, color .15s;
  }
  .iconbtn:not(:disabled):hover { border-color: var(--accent); color: var(--accent); }
  .iconbtn:disabled { opacity: .3; cursor: not-allowed; }
  .iconbtn.del:not(:disabled):hover { border-color: var(--danger); color: var(--danger); }

  .actions { display: flex; align-items: center; gap: 10px; margin-top: 20px; flex-wrap: wrap; }
  .status { font-family: var(--mono); font-size: 11px; color: var(--muted); margin-left: auto; }
  .status.ok { color: var(--success); }
  .status.err { color: var(--danger); }

  .splitmeta { font-family: var(--mono); font-size: 11px; color: var(--muted); margin-top: 14px; }
  .toolrow { display: flex; align-items: center; gap: 10px; margin-top: 16px; flex-wrap: wrap; }
  .rangelbl { font-family: var(--mono); font-size: 11px; color: var(--muted); display: flex; align-items: center; gap: 8px; }
  .rangebar {
    background: var(--field); border: 1px solid var(--line-strong); border-radius: var(--r);
    color: var(--ink); font-family: var(--mono); font-size: 12px; padding: 8px 10px;
    width: 200px; outline: none; transition: border-color .15s;
  }
  .rangebar:focus { border-color: var(--accent); }

  .grid { display: grid; grid-template-columns: repeat(auto-fill, minmax(104px, 1fr)); gap: 12px; margin-top: 18px; }
  .cell { cursor: pointer; text-align: center; user-select: none; }
  .pagebox {
    position: relative; border: 1px solid var(--line-strong); border-radius: var(--r);
    overflow: hidden; background: var(--paper-2); aspect-ratio: 3 / 4;
    display: flex; align-items: center; justify-content: center;
    color: var(--muted); font-family: var(--mono); font-size: 11px;
    transition: border-color .15s, opacity .15s;
  }
  .cell.sel .pagebox { border-color: var(--accent); }
  .cell:not(.sel) .pagebox { opacity: .4; }
  .pagebox img { width: 100%; height: 100%; object-fit: contain; background: #fff; }
  .check {
    position: absolute; top: 4px; right: 4px; width: 18px; height: 18px; border-radius: var(--r);
    background: var(--accent); color: var(--on-accent); font-size: 11px; font-weight: 700;
    display: none; align-items: center; justify-content: center;
  }
  .cell.sel .check { display: flex; }
  .plabel { font-family: var(--mono); font-size: 10px; color: var(--muted); margin-top: 6px; }

  .levels { display: flex; gap: 6px; flex-wrap: wrap; }
  .level {
    background: transparent; border: 1px solid var(--line-strong); color: var(--muted);
    padding: 7px 14px; border-radius: var(--r); cursor: pointer; font-family: var(--mono);
    font-size: 11px; font-weight: 700; letter-spacing: 1px; text-transform: uppercase;
    transition: border-color .15s, color .15s, background .15s;
  }
  .level:hover { color: var(--ink); border-color: var(--accent); }
  .level.active { background: var(--accent); color: var(--on-accent); border-color: var(--accent); }
  .cinfo { display: flex; gap: 14px; align-items: center; margin-top: 18px; }
  .cthumb {
    width: 56px; height: 74px; border-radius: var(--r); border: 1px solid var(--line);
    object-fit: cover; object-position: top center; background: var(--paper-2); flex: 0 0 auto;
  }
  .cname { font-family: var(--mono); font-size: 12px; color: var(--ink); }
  .cresult {
    margin-top: 18px; font-family: var(--mono); font-size: 12px; color: var(--ink-soft);
    line-height: 1.8; background: var(--paper-2); border: 1px solid var(--line);
    border-radius: var(--r); padding: 12px 14px;
  }
  .cresult .save { color: var(--success); font-weight: 700; }
  .cresult .grow { color: var(--danger); font-weight: 700; }

  /* Page preview */
  .modal {
    position: fixed; inset: 0; background: rgba(0,0,0,.90); z-index: 50; overflow: auto;
    display: flex; flex-direction: column; align-items: center; padding: 24px 16px;
  }
  .modal.hidden { display: none; }
  .modal-bar { width: 100%; max-width: 820px; display: flex; align-items: center; gap: 12px; margin-bottom: 14px; }
  .modal-bar .title {
    font-family: var(--mono); font-size: 12px; color: var(--ink); flex: 1 1 auto;
    white-space: nowrap; overflow: hidden; text-overflow: ellipsis;
  }
  .modal-close {
    background: transparent; border: 1px solid var(--line-strong); color: var(--ink);
    width: 32px; height: 32px; border-radius: var(--r); cursor: pointer; font-size: 13px;
    transition: border-color .15s, color .15s;
  }
  .modal-close:hover { border-color: var(--sky); color: var(--sky); }
  .pages { display: flex; flex-direction: column; gap: 14px; align-items: center; }
  .pages canvas { max-width: 100%; background: #fff; border-radius: var(--r); border: 1px solid var(--line); }
</style>
</head>
<body>
<div class="container">

  <header>
    <div class="logo-frame">
      <div class="logo-word">MARKDOWN</div>
      <div class="logo-sub"><span class="logo-rule"></span>FACTORY</div>
    </div>
    <div class="tagline">Every document becomes Markdown — all of it on this computer.</div>
    <div class="version">V__VERSION__</div>
  </header>

  <div class="tabs">
    <button class="tab active" data-tab="docs">Documents</button>
    <button class="tab" data-tab="ocr">Scans &amp; Photos</button>
    <button class="tab" data-tab="pdf">PDF Tools</button>
    <button class="tab" data-tab="word">To Word</button>
  </div>

  <div class="outrow" id="outrow">
    <label for="outdir">SAVE TO</label>
    <input id="outdir" value="__DEFAULT_OUTPUT__">
    <button id="openBtn" class="ghost btn-sm" title="Open output folder">Open</button>
  </div>

  <div class="tab-content active" id="tab-docs">
    <p class="sub">Word, PowerPoint, Excel, HTML and PDFs that already contain real text — converted on this machine, each file coming back as one self-contained <code>.md</code>. PDFs get a purpose-built engine that keeps headings, tables and links; a PDF that turns out to be a scan says so and points you at Scans &amp; Photos, rather than quietly converting into an empty file.</p>
    <div class="card">
      <div id="drop">
        <p class="big">Drag &amp; drop files here</p>
        <p class="small">or click to browse</p>
        <input type="file" id="fileInput" multiple>
      </div>
      <ul id="list"></ul>
      <div id="summary"></div>
    </div>
  </div>

  <div class="tab-content" id="tab-ocr">
    <div class="card" id="ocrCard">
      <h2 id="ocrTitle">Checking this computer…</h2>
      <p id="ocrText"></p>
      <button id="setupBtn" class="go" style="display:none">Set up OCR (one-time)</button>
      <div id="setupStage"></div>
      <div id="setupLog"></div>
    </div>
    <iframe id="ocrFrame" src="about:blank"></iframe>
  </div>

  <div class="tab-content" id="tab-pdf">
    <div class="subtabs">
      <button id="pdfTabMerge" class="subtab active">Merge</button>
      <button id="pdfTabSplit" class="subtab">Split</button>
      <button id="pdfTabCompress" class="subtab">Compress</button>
    </div>
    <div id="pdfBoot"></div>

    <section id="mergeView">
      <p class="sub">Drag in PDFs, reorder them, and merge into a single file. Click a thumbnail to flip through a file before you commit.</p>
      <div class="card">
        <div id="dropzone" class="dropzone">
          <div class="big">Drop PDF files here</div>
          <div class="small">or click to browse</div>
          <input id="pdfFileInput" type="file" accept="application/pdf,.pdf" multiple class="hidden">
        </div>
        <ul id="pdfList" class="list"></ul>
        <div class="actions">
          <button id="mergeBtn" class="go" disabled>Merge &amp; save</button>
          <button id="clearBtn" class="ghost hidden">Clear all</button>
          <span id="status" class="status"></span>
        </div>
      </div>
    </section>

    <section id="splitView" class="hidden">
      <p class="sub">Load one PDF, pick the pages you want, then extract them into a new file — or split every page into its own file.</p>
      <div class="card">
        <div id="splitDrop" class="dropzone">
          <div class="big">Drop a PDF here</div>
          <div class="small">or click to browse</div>
          <input id="splitInput" type="file" accept="application/pdf,.pdf" class="hidden">
        </div>
        <div id="splitMeta" class="splitmeta"></div>
        <div id="splitTools" class="toolrow hidden">
          <label class="rangelbl">PAGES
            <input id="splitRange" class="rangebar" type="text" placeholder="e.g. 1-3, 5, 8-10">
          </label>
          <button id="selectAllBtn" class="ghost btn-sm">Select all</button>
          <button id="selectNoneBtn" class="ghost btn-sm">Select none</button>
        </div>
        <div id="splitGrid" class="grid"></div>
        <div class="actions">
          <button id="extractBtn" class="go" disabled>Extract selected</button>
          <button id="splitAllBtn" class="ghost" disabled>Split each page</button>
          <button id="splitClearBtn" class="ghost hidden">Clear</button>
          <span id="splitStatus" class="status"></span>
        </div>
      </div>
    </section>

    <section id="compressView" class="hidden">
      <p class="sub">Shrink a PDF by re-rendering its pages as compressed images — best for scans and image-heavy files. This flattens pages to images, so text stops being selectable.</p>
      <div class="card">
        <div id="compressDrop" class="dropzone">
          <div class="big">Drop a PDF here</div>
          <div class="small">or click to browse</div>
          <input id="compressInput" type="file" accept="application/pdf,.pdf" class="hidden">
        </div>
        <div id="compressPanel" class="hidden">
          <div class="cinfo">
            <img id="compressThumb" class="cthumb" alt="preview">
            <div>
              <div id="compressName" class="cname"></div>
              <div id="compressMeta" class="info"></div>
            </div>
          </div>
          <div class="toolrow">
            <span class="rangelbl">LEVEL</span>
            <div id="levelBtns" class="levels">
              <button class="level" data-level="light">Light</button>
              <button class="level active" data-level="balanced">Balanced</button>
              <button class="level" data-level="strong">Strong</button>
              <button class="level" data-level="extreme">Extreme</button>
            </div>
          </div>
          <div class="actions">
            <button id="compressBtn" class="go">Compress</button>
            <button id="compressSaveBtn" class="ghost hidden">Save compressed PDF</button>
            <button id="compressClearBtn" class="ghost hidden">Clear</button>
            <span id="compressStatus" class="status"></span>
          </div>
          <div id="compressResult" class="cresult hidden"></div>
        </div>
      </div>
    </section>

    <div id="modal" class="modal hidden">
      <div class="modal-bar">
        <span id="modalTitle" class="title"></span>
        <button id="modalClose" class="modal-close" title="Close">✕</button>
      </div>
      <div id="pages" class="pages"></div>
    </div>
  </div>

  <div class="tab-content" id="tab-word">
    <p class="sub">The way back out. <strong>PDFs</strong> that already contain real text are rebuilt as editable Word files, with headings, tables, columns and pictures kept where they were — a scanned PDF has no text to rebuild, so read it on the <strong>Scans &amp; Photos</strong> tab instead. <strong>Markdown</strong> files become properly styled Word documents: real headings, lists, tables, quotes and links, with code in a monospaced block. One <code>.docx</code> per file either way.</p>
    <div class="card">
      <div id="wordDrop" class="dropzone">
        <div class="big">Drop PDF or Markdown files here</div>
        <div class="small">or click to browse</div>
        <input id="wordInput" type="file" accept="application/pdf,.pdf,text/markdown,.md,.markdown,.mdown,.mkd" multiple class="hidden">
      </div>
      <div class="toolrow">
        <label class="rangelbl">PAGES <span class="hint">(PDFs only)</span>
          <input id="wordRange" class="rangebar" type="text" placeholder="all pages — or 1-3, 5, 8-10">
        </label>
      </div>
      <ul id="wordList"></ul>
      <div id="wordSummary"></div>
    </div>
  </div>
</div>
<script>
const outrow = document.getElementById('outrow');
document.querySelectorAll('.tab').forEach(btn => {
  btn.addEventListener('click', () => {
    document.querySelectorAll('.tab').forEach(b => b.classList.remove('active'));
    document.querySelectorAll('.tab-content').forEach(c => c.classList.remove('active'));
    btn.classList.add('active');
    document.getElementById('tab-' + btn.dataset.tab).classList.add('active');
    // the OCR pane carries its own output controls inside the worker page
    outrow.style.display = btn.dataset.tab === 'ocr' ? 'none' : 'flex';
    if (btn.dataset.tab === 'pdf' && window.bootPdfTools) window.bootPdfTools();
  });
});

// dropping a file outside a drop zone should never navigate the window away
window.addEventListener('dragover', e => e.preventDefault());
window.addEventListener('drop', e => e.preventDefault());

function bytesToB64(bytes) {
  let bin = '';
  const CHUNK = 0x8000;
  for (let i = 0; i < bytes.length; i += CHUNK)
    bin += String.fromCharCode.apply(null, bytes.subarray(i, i + CHUNK));
  return btoa(bin);
}

function baseName(p) { return p.split(/[\\\\/]/).pop(); }

/* ---- Documents tab (markitdown) ---- */
const drop = document.getElementById('drop');
const fileInput = document.getElementById('fileInput');
const list = document.getElementById('list');
const summary = document.getElementById('summary');
const outdir = document.getElementById('outdir');
let total = 0, done = 0, failed = 0;
let queue = Promise.resolve();

drop.addEventListener('click', () => fileInput.click());
fileInput.addEventListener('change', () => { handleFiles(fileInput.files); fileInput.value = ''; });
['dragenter','dragover'].forEach(ev => drop.addEventListener(ev, e => {
  e.preventDefault(); drop.classList.add('over');
}));
['dragleave','drop'].forEach(ev => drop.addEventListener(ev, e => {
  e.preventDefault(); drop.classList.remove('over');
}));
drop.addEventListener('drop', e => handleFiles(e.dataTransfer.files));

function handleFiles(files) {
  for (const f of files) {
    total++;
    const li = document.createElement('li');
    li.innerHTML = `<div class="top"><span class="name"></span><span class="st wait">queued</span></div>` +
                   `<div class="note"></div>`;
    li.querySelector('.name').textContent = f.name;   // textContent, not innerHTML:
    li.querySelector('.name').title = f.name;         // a filename may contain markup
    list.prepend(li);
    queue = queue.then(() => convertFile(f, li));
  }
  updateSummary();
}

async function convertFile(file, li) {
  const st = li.querySelector('.st');
  const note = li.querySelector('.note');
  st.textContent = 'converting…'; st.className = 'st run';
  try {
    const bytes = new Uint8Array(await file.arrayBuffer());
    const resp = await fetch('/convert-doc', {
      method: 'POST', headers: {'Content-Type': 'application/json'},
      body: JSON.stringify({ name: file.name, data: bytesToB64(bytes), outdir: outdir.value })
    });
    const r = await resp.json();
    if (r.ok) {
      st.textContent = 'done'; st.className = 'st ok'; st.title = r.path;
      note.textContent = r.path; note.classList.remove('failed'); done++;
    } else {
      st.textContent = 'failed'; st.className = 'st err'; st.title = r.error;
      note.textContent = r.error; note.classList.add('failed'); failed++;
    }
  } catch (e) {
    st.textContent = 'failed'; st.className = 'st err'; st.title = String(e);
    note.textContent = String(e.message || e); note.classList.add('failed'); failed++;
  }
  updateSummary();
}

function updateSummary() {
  summary.textContent = total ? `${done} converted, ${failed} failed, ${total - done - failed} pending` : '';
}

document.getElementById('openBtn').addEventListener('click', () =>
  fetch('/open', { method: 'POST', headers: {'Content-Type':'application/json'},
                   body: JSON.stringify({ outdir: outdir.value }) }));

/* ---- To Word tab (pdf2docx / markdown-it) ----
   PDFs cost seconds per page rather than seconds per file, so the POST only
   starts a job and hands back an id; the row then polls it and shows which page
   is being read or written. Files convert one at a time — the engine is
   single-threaded and a queue that pretends otherwise just makes every row
   slower. Markdown takes the same road for one shape of row and one code path,
   and simply arrives at 'done' on the first poll. */
(function () {
  const wdrop = document.getElementById('wordDrop');
  const winput = document.getElementById('wordInput');
  const wlist = document.getElementById('wordList');
  const wrange = document.getElementById('wordRange');
  const wsummary = document.getElementById('wordSummary');
  let wtotal = 0, wdone = 0, wfailed = 0;
  let wqueue = Promise.resolve();

  wdrop.addEventListener('click', () => winput.click());
  winput.addEventListener('change', () => { addFiles(winput.files); winput.value = ''; });
  ['dragenter','dragover'].forEach(ev => wdrop.addEventListener(ev, e => {
    e.preventDefault(); wdrop.classList.add('drag');
  }));
  ['dragleave','drop'].forEach(ev => wdrop.addEventListener(ev, e => {
    e.preventDefault(); wdrop.classList.remove('drag');
  }));
  wdrop.addEventListener('drop', e => addFiles(e.dataTransfer.files));

  function addFiles(files) {
    for (const f of files) {
      wtotal++;
      const li = document.createElement('li');
      li.innerHTML = `<div class="top"><span class="name"></span><span class="st wait">queued</span></div>` +
                     `<div class="bar"><i></i></div><div class="note">waiting…</div>`;
      li.querySelector('.name').textContent = f.name;
      li.querySelector('.name').title = f.name;
      wlist.prepend(li);
      // the range is read when the file's turn comes, so editing it mid-queue
      // would silently change earlier rows — freeze it at drop time
      const pages = wrange.value;
      wqueue = wqueue.then(() => convertToWord(f, pages, li));
    }
    updateWordSummary();
  }

  async function convertToWord(file, pages, li) {
    const st = li.querySelector('.st');
    const bar = li.querySelector('.bar');
    const fill = li.querySelector('.bar > i');
    const note = li.querySelector('.note');
    st.textContent = 'converting…'; st.className = 'st run';
    note.textContent = 'uploading…';
    try {
      const bytes = new Uint8Array(await file.arrayBuffer());
      const resp = await fetch('/convert-word', {
        method: 'POST', headers: {'Content-Type': 'application/json'},
        body: JSON.stringify({ name: file.name, data: bytesToB64(bytes),
                               outdir: outdir.value, pages: pages })
      });
      const started = await resp.json();
      if (!started.ok) throw new Error(started.error);

      for (;;) {
        await new Promise(r => setTimeout(r, 700));
        const r = await (await fetch('/word/state?job=' + started.job)).json();
        if (!r.ok) throw new Error(r.error || 'lost track of this conversion');
        fill.style.width = r.pct + '%';
        if (r.state === 'done') {
          st.textContent = 'done'; st.className = 'st ok';
          bar.classList.add('done'); fill.style.width = '100%';
          note.textContent = r.path;
          wdone++;
          break;
        }
        if (r.state === 'error') throw new Error(r.error);
        note.textContent = r.note;
      }
    } catch (e) {
      st.textContent = 'failed'; st.className = 'st err';
      bar.classList.add('failed'); fill.style.width = '100%';
      note.textContent = String(e.message || e); note.classList.add('failed');
      wfailed++;
    }
    updateWordSummary();
  }

  function updateWordSummary() {
    wsummary.textContent = wtotal
      ? `${wdone} converted, ${wfailed} failed, ${wtotal - wdone - wfailed} pending` : '';
  }
})();

/* ---- OCR tab (managed worker) ---- */
const ocrCard = document.getElementById('ocrCard');
const ocrTitle = document.getElementById('ocrTitle');
const ocrText = document.getElementById('ocrText');
const setupBtn = document.getElementById('setupBtn');
const setupStage = document.getElementById('setupStage');
const setupLog = document.getElementById('setupLog');
const ocrFrame = document.getElementById('ocrFrame');

setupBtn.addEventListener('click', async () => {
  setupBtn.disabled = true;
  await fetch('/ocr/setup', { method: 'POST' });
});

async function pollOcr() {
  try {
    const r = await (await fetch('/ocr/state')).json();
    if (r.state === 'ready') {
      ocrCard.style.display = 'none';
      if (ocrFrame.src === 'about:blank') ocrFrame.src = 'http://127.0.0.1:__WORKER_PORT__/';
      ocrFrame.style.display = 'block';
      return; // engine page takes over from here
    }
    ocrFrame.style.display = 'none';
    ocrCard.style.display = 'block';
    setupBtn.style.display = 'none';
    setupLog.style.display = 'none';
    setupStage.textContent = '';
    if (r.state === 'no-gpu') {
      ocrTitle.textContent = 'This computer can\\'t run the OCR engine';
      ocrText.textContent = 'Reading scans and photos needs an NVIDIA graphics card, and none was found. ' +
                            'The Documents tab still works normally.';
    } else if (r.state === 'needs-setup') {
      ocrTitle.textContent = 'One-time setup needed';
      ocrText.textContent = 'The OCR engine downloads about 4 GB of libraries the first time (plus a 6 GB ' +
                            'reading model on first use). Needs an NVIDIA graphics card and an internet connection.';
      setupBtn.style.display = 'inline-block';
      setupBtn.disabled = false;
    } else if (r.state === 'installing') {
      ocrTitle.textContent = 'Setting up the OCR engine…';
      ocrText.textContent = 'This is a one-time download of a few gigabytes — a good moment for coffee. ' +
                            'Keep this window open.';
      setupStage.textContent = r.detail || '';
      setupLog.style.display = 'block';
      setupLog.textContent = (r.log || []).join('\\n');
      setupLog.scrollTop = setupLog.scrollHeight;
    } else if (r.state === 'starting') {
      ocrTitle.textContent = 'Starting the OCR engine…';
      ocrText.textContent = '';
      setupStage.textContent = r.detail || '';
    } else if (r.state === 'error') {
      ocrTitle.textContent = 'OCR setup hit a problem';
      ocrText.textContent = r.detail || 'See the log below.';
      setupBtn.textContent = 'Try again';
      setupBtn.style.display = 'inline-block';
      setupBtn.disabled = false;
      setupLog.style.display = 'block';
      setupLog.textContent = (r.log || []).join('\\n');
    } else {
      ocrTitle.textContent = 'Checking this computer…';
    }
  } catch (e) {}
  setTimeout(pollOcr, 1500);
}
pollOcr();

/* ---- PDF Tools tab (merge / split / compress) ----
   Absorbed from the standalone PDF Tools app. pdf-lib + pdf.js are ~1.9 MB, so
   they load from /vendor the first time this tab is opened rather than on every
   page load. Everything still runs in the browser; only the finished PDF is
   posted back, so it lands in the shared export folder. */
(function () {
  const boot = document.getElementById('pdfBoot');
  let libsPromise = null;

  function loadScript(src) {
    return new Promise((resolve, reject) => {
      const s = document.createElement('script');
      s.src = src;
      s.onload = resolve;
      s.onerror = () => reject(new Error('could not load ' + src));
      document.head.appendChild(s);
    });
  }

  function loadLibs() {
    if (!libsPromise) {
      boot.textContent = 'Loading the PDF engine…';
      boot.className = '';
      libsPromise = (async () => {
        await loadScript('/vendor/pdf-lib.min.js');
        await loadScript('/vendor/pdf.min.js');
        pdfjsLib.GlobalWorkerOptions.workerSrc = '/vendor/pdf.worker.min.js';
      })();
      libsPromise.then(
        () => { boot.textContent = ''; },
        () => { boot.textContent = 'The PDF engine failed to load — vendor files missing?'; boot.className = 'err'; }
      );
    }
    return libsPromise;
  }
  window.bootPdfTools = loadLibs;   // the tab button warms it up on first click

  async function ready(setter) {
    try { await loadLibs(); return true; }
    catch (e) { setter('The PDF engine is not available.', 'err'); return false; }
  }

  /* the finished bytes go to the server so PDFs land beside the Markdown */
  async function saveToFolder(bytes, name) {
    const resp = await fetch('/save-pdf', {
      method: 'POST', headers: { 'Content-Type': 'application/json' },
      body: JSON.stringify({ name, outdir: outdir.value, data: bytesToB64(bytes) })
    });
    const r = await resp.json();
    if (!r.ok) throw new Error(r.error);
    return r.path;
  }

  function fmtSize(b) {
    if (b < 1024) return b + ' B';
    if (b < 1024 * 1024) return (b / 1024).toFixed(0) + ' KB';
    return (b / 1024 / 1024).toFixed(1) + ' MB';
  }

  async function renderPageToDataURL(pdf, pageNum, scale) {
    const page = await pdf.getPage(pageNum);
    const viewport = page.getViewport({ scale });
    const canvas = document.createElement('canvas');
    canvas.width = Math.ceil(viewport.width);
    canvas.height = Math.ceil(viewport.height);
    const ctx = canvas.getContext('2d');
    ctx.fillStyle = '#fff';
    ctx.fillRect(0, 0, canvas.width, canvas.height);
    await page.render({ canvasContext: ctx, viewport }).promise;
    return canvas.toDataURL('image/jpeg', 0.7);
  }

  /* ---- sub-tabs ---- */
  const TABS = {
    merge:    { tab: document.getElementById('pdfTabMerge'),    view: document.getElementById('mergeView') },
    split:    { tab: document.getElementById('pdfTabSplit'),    view: document.getElementById('splitView') },
    compress: { tab: document.getElementById('pdfTabCompress'), view: document.getElementById('compressView') },
  };
  for (const key in TABS) {
    TABS[key].tab.addEventListener('click', () => {
      for (const k in TABS) {
        TABS[k].view.classList.toggle('hidden', k !== key);
        TABS[k].tab.classList.toggle('active', k === key);
      }
    });
  }

  /* ===================== Merge ===================== */
  const dropzone = document.getElementById('dropzone');
  const pdfFileInput = document.getElementById('pdfFileInput');
  const listEl = document.getElementById('pdfList');
  const mergeBtn = document.getElementById('mergeBtn');
  const clearBtn = document.getElementById('clearBtn');
  const statusEl = document.getElementById('status');
  const modal = document.getElementById('modal');
  const modalTitle = document.getElementById('modalTitle');
  const modalClose = document.getElementById('modalClose');
  const pagesEl = document.getElementById('pages');

  let files = []; // { id, name, size, bytes, pages, thumb }
  let nextId = 1;

  function setStatus(msg, cls = '') {
    statusEl.textContent = msg;
    statusEl.className = 'status ' + cls;
  }

  async function addFiles(fileList) {
    const pdfs = Array.from(fileList).filter(f =>
      f.type === 'application/pdf' || f.name.toLowerCase().endsWith('.pdf'));
    if (pdfs.length === 0) {
      setStatus('No PDF files found in that drop.', 'err');
      return;
    }
    setStatus('Loading ' + pdfs.length + ' file(s)…');
    if (!await ready(setStatus)) return;
    const added = [];
    for (const f of pdfs) {
      try {
        const bytes = new Uint8Array(await f.arrayBuffer());
        const entry = { id: nextId++, name: f.name, size: f.size, bytes, pages: null, thumb: null };
        files.push(entry);
        added.push(entry);
      } catch (e) {
        setStatus('Could not read "' + f.name + '".', 'err');
      }
    }
    render();
    setStatus('');
    // Render thumbnails + page counts asynchronously so the list shows immediately.
    for (const entry of added) {
      try {
        const pdf = await pdfjsLib.getDocument({ data: entry.bytes.slice(0), disableAutoFetch: true }).promise;
        entry.pages = pdf.numPages;
        entry.thumb = await renderPageToDataURL(pdf, 1, 0.4);
        await pdf.destroy();
      } catch (e) {
        // Encrypted/odd files: fall back to pdf-lib for the page count, leave thumb empty.
        try {
          const doc = await PDFLib.PDFDocument.load(entry.bytes, { ignoreEncryption: true });
          entry.pages = doc.getPageCount();
        } catch (e2) { entry.pages = null; }
      }
      render();
    }
  }

  function render() {
    listEl.innerHTML = '';
    files.forEach((f, i) => {
      const li = document.createElement('li');
      li.className = 'item';
      li.draggable = true;
      li.dataset.id = f.id;

      const pageInfo = f.pages == null ? 'page count unknown' :
        (f.pages + (f.pages === 1 ? ' page' : ' pages'));

      const thumbInner = f.thumb
        ? `<img src="${f.thumb}" alt="preview">`
        : (f.pages === null ? '…' : '📄');

      li.innerHTML = `
        <span class="grip">⋮⋮</span>
        <span class="idx">${i + 1}</span>
        <span class="thumb${f.thumb ? '' : ' loading'}" title="Click to preview all pages">${thumbInner}</span>
        <span class="meta">
          <div class="name" title="${f.name.replace(/"/g, '&quot;')}">${f.name}</div>
          <div class="info">${pageInfo} · ${fmtSize(f.size)}</div>
        </span>
        <span class="btns">
          <button class="iconbtn up" ${i === 0 ? 'disabled' : ''} title="Move up">↑</button>
          <button class="iconbtn down" ${i === files.length - 1 ? 'disabled' : ''} title="Move down">↓</button>
          <button class="iconbtn del" title="Remove">✕</button>
        </span>`;

      li.querySelector('.up').onclick = () => move(i, i - 1);
      li.querySelector('.down').onclick = () => move(i, i + 1);
      li.querySelector('.del').onclick = () => { files.splice(i, 1); render(); };
      const thumbEl = li.querySelector('.thumb');
      thumbEl.draggable = false;
      thumbEl.onclick = (e) => { e.stopPropagation(); openPreview(f); };

      // drag to reorder
      li.addEventListener('dragstart', e => {
        li.classList.add('dragging');
        e.dataTransfer.effectAllowed = 'move';
        e.dataTransfer.setData('text/plain', String(f.id));
      });
      li.addEventListener('dragend', () => {
        li.classList.remove('dragging');
        listEl.querySelectorAll('.item.over').forEach(el => el.classList.remove('over'));
      });
      li.addEventListener('dragover', e => {
        e.preventDefault();
        e.stopPropagation();
        li.classList.add('over');
      });
      li.addEventListener('dragleave', () => li.classList.remove('over'));
      li.addEventListener('drop', e => {
        e.preventDefault();
        e.stopPropagation();
        li.classList.remove('over');
        const draggedId = Number(e.dataTransfer.getData('text/plain'));
        const from = files.findIndex(x => x.id === draggedId);
        const to = files.findIndex(x => x.id === f.id);
        if (from === -1 || to === -1 || from === to) return;
        const [moved] = files.splice(from, 1);
        files.splice(to, 0, moved);
        render();
      });

      listEl.appendChild(li);
    });

    const has = files.length > 0;
    mergeBtn.disabled = !has;
    clearBtn.classList.toggle('hidden', !has);
  }

  function move(from, to) {
    if (to < 0 || to >= files.length) return;
    const [m] = files.splice(from, 1);
    files.splice(to, 0, m);
    render();
  }

  async function openPreview(f) {
    modalTitle.textContent = f.name;
    pagesEl.innerHTML = '<div style="color:var(--muted)">Rendering pages…</div>';
    modal.classList.remove('hidden');
    try {
      const pdf = await pdfjsLib.getDocument({ data: f.bytes.slice(0) }).promise;
      pagesEl.innerHTML = '';
      const scale = Math.min(1.5, Math.max(0.8, (Math.min(800, window.innerWidth - 60)) / 612));
      for (let n = 1; n <= pdf.numPages; n++) {
        if (modal.classList.contains('hidden')) break; // user closed mid-render
        const page = await pdf.getPage(n);
        const viewport = page.getViewport({ scale });
        const canvas = document.createElement('canvas');
        canvas.width = Math.ceil(viewport.width);
        canvas.height = Math.ceil(viewport.height);
        await page.render({ canvasContext: canvas.getContext('2d'), viewport }).promise;
        pagesEl.appendChild(canvas);
      }
      await pdf.destroy();
    } catch (e) {
      pagesEl.innerHTML = '<div style="color:var(--danger)">Could not render this PDF (it may be encrypted).</div>';
    }
  }

  function closePreview() {
    modal.classList.add('hidden');
    pagesEl.innerHTML = '';
  }

  modalClose.addEventListener('click', closePreview);
  modal.addEventListener('click', e => { if (e.target === modal) closePreview(); });
  document.addEventListener('keydown', e => { if (e.key === 'Escape') closePreview(); });

  async function mergePdfs() {
    if (files.length === 0) return;
    mergeBtn.disabled = true;
    setStatus('Merging…');
    try {
      const out = await PDFLib.PDFDocument.create();
      for (const f of files) {
        const src = await PDFLib.PDFDocument.load(f.bytes, { ignoreEncryption: true });
        const copied = await out.copyPages(src, src.getPageIndices());
        copied.forEach(p => out.addPage(p));
      }
      const path = await saveToFolder(await out.save(), 'merged.pdf');
      setStatus('Merged ' + files.length + ' files → ' + baseName(path), 'ok');
    } catch (e) {
      console.error(e);
      setStatus('Merge failed: ' + e.message, 'err');
    } finally {
      mergeBtn.disabled = files.length < 1;
    }
  }

  dropzone.addEventListener('click', () => pdfFileInput.click());
  pdfFileInput.addEventListener('change', () => { addFiles(pdfFileInput.files); pdfFileInput.value = ''; });
  ['dragenter', 'dragover'].forEach(ev =>
    dropzone.addEventListener(ev, e => { e.preventDefault(); dropzone.classList.add('drag'); }));
  ['dragleave', 'drop'].forEach(ev =>
    dropzone.addEventListener(ev, e => { e.preventDefault(); dropzone.classList.remove('drag'); }));
  dropzone.addEventListener('drop', e => {
    e.stopPropagation();
    if (e.dataTransfer.files && e.dataTransfer.files.length) addFiles(e.dataTransfer.files);
  });
  mergeBtn.addEventListener('click', mergePdfs);
  clearBtn.addEventListener('click', () => { files = []; render(); setStatus(''); });

  /* ===================== Split ===================== */
  const splitDrop = document.getElementById('splitDrop');
  const splitInput = document.getElementById('splitInput');
  const splitMeta = document.getElementById('splitMeta');
  const splitTools = document.getElementById('splitTools');
  const splitRange = document.getElementById('splitRange');
  const splitGrid = document.getElementById('splitGrid');
  const extractBtn = document.getElementById('extractBtn');
  const splitAllBtn = document.getElementById('splitAllBtn');
  const splitClearBtn = document.getElementById('splitClearBtn');
  const selectAllBtn = document.getElementById('selectAllBtn');
  const selectNoneBtn = document.getElementById('selectNoneBtn');
  const splitStatus = document.getElementById('splitStatus');

  let splitDoc = null; // { name, bytes, pages, selected:Set<number> }

  function setSplitStatus(msg, cls = '') {
    splitStatus.textContent = msg;
    splitStatus.className = 'status ' + cls;
  }

  function splitBaseName() { return splitDoc.name.replace(/\\.pdf$/i, ''); }

  async function loadSplitFile(file) {
    if (!file || !(file.type === 'application/pdf' || file.name.toLowerCase().endsWith('.pdf'))) {
      setSplitStatus('Please choose a PDF file.', 'err');
      return;
    }
    setSplitStatus('Loading…');
    if (!await ready(setSplitStatus)) return;
    try {
      const bytes = new Uint8Array(await file.arrayBuffer());
      const pdf = await pdfjsLib.getDocument({ data: bytes.slice(0) }).promise;
      splitDoc = { name: file.name, bytes, pages: pdf.numPages, selected: new Set() };
      for (let i = 1; i <= pdf.numPages; i++) splitDoc.selected.add(i); // select all by default
      renderSplit();
      setSplitStatus('');
      // fill in thumbnails progressively
      for (let n = 1; n <= pdf.numPages; n++) {
        const url = await renderPageToDataURL(pdf, n, 0.5);
        const img = splitGrid.querySelector('[data-page="' + n + '"] img');
        if (img) img.src = url;
      }
      await pdf.destroy();
    } catch (e) {
      console.error(e);
      setSplitStatus('Could not read this PDF (it may be encrypted).', 'err');
      splitDoc = null;
      renderSplit();
    }
  }

  function renderSplit() {
    if (!splitDoc) {
      splitMeta.textContent = '';
      splitGrid.innerHTML = '';
      splitTools.classList.add('hidden');
      updateSplitButtons();
      return;
    }
    splitMeta.textContent = splitDoc.name + ' · ' + splitDoc.pages +
      (splitDoc.pages === 1 ? ' page' : ' pages');
    splitTools.classList.remove('hidden');
    splitGrid.innerHTML = '';
    for (let n = 1; n <= splitDoc.pages; n++) {
      const cell = document.createElement('div');
      cell.className = 'cell' + (splitDoc.selected.has(n) ? ' sel' : '');
      cell.dataset.page = n;
      cell.innerHTML =
        '<div class="pagebox"><img alt="p' + n + '"><span class="check">✓</span></div>' +
        '<div class="plabel">' + n + '</div>';
      cell.addEventListener('click', () => {
        if (splitDoc.selected.has(n)) splitDoc.selected.delete(n);
        else splitDoc.selected.add(n);
        cell.classList.toggle('sel');
        syncRangeFromSelection();
        updateSplitButtons();
      });
      splitGrid.appendChild(cell);
    }
    syncRangeFromSelection();
    updateSplitButtons();
  }

  function updateSplitButtons() {
    const loaded = !!splitDoc;
    const hasSel = loaded && splitDoc.selected.size > 0;
    extractBtn.disabled = !hasSel;
    splitAllBtn.disabled = !loaded;
    splitClearBtn.classList.toggle('hidden', !loaded);
  }

  function selectionToRangeStr() {
    const arr = [...splitDoc.selected].sort((a, b) => a - b);
    const parts = [];
    let start = null, prev = null;
    for (const n of arr) {
      if (start === null) { start = n; prev = n; }
      else if (n === prev + 1) { prev = n; }
      else { parts.push(start === prev ? '' + start : start + '-' + prev); start = n; prev = n; }
    }
    if (start !== null) parts.push(start === prev ? '' + start : start + '-' + prev);
    return parts.join(', ');
  }

  function syncRangeFromSelection() { splitRange.value = selectionToRangeStr(); }

  function parseRange(str, max) {
    const set = new Set();
    for (let token of str.split(',')) {
      token = token.trim();
      if (!token) continue;
      const m = token.match(/^(\\d+)\\s*-\\s*(\\d+)$/);
      if (m) {
        let a = +m[1], b = +m[2];
        if (a > b) { const t = a; a = b; b = t; }
        for (let i = a; i <= b; i++) if (i >= 1 && i <= max) set.add(i);
      } else if (/^\\d+$/.test(token)) {
        const v = +token;
        if (v >= 1 && v <= max) set.add(v);
      }
    }
    return set;
  }

  function applyRangeInput() {
    if (!splitDoc) return;
    splitDoc.selected = parseRange(splitRange.value, splitDoc.pages);
    splitGrid.querySelectorAll('.cell').forEach(c => {
      c.classList.toggle('sel', splitDoc.selected.has(+c.dataset.page));
    });
    syncRangeFromSelection();
    updateSplitButtons();
  }

  async function extractSelected() {
    if (!splitDoc || splitDoc.selected.size === 0) return;
    extractBtn.disabled = true;
    setSplitStatus('Extracting…');
    try {
      const src = await PDFLib.PDFDocument.load(splitDoc.bytes, { ignoreEncryption: true });
      const out = await PDFLib.PDFDocument.create();
      const idx = [...splitDoc.selected].sort((a, b) => a - b).map(n => n - 1);
      const copied = await out.copyPages(src, idx);
      copied.forEach(p => out.addPage(p));
      const path = await saveToFolder(await out.save(),
        splitBaseName() + ' (pages ' + selectionToRangeStr() + ').pdf');
      setSplitStatus('Extracted ' + idx.length + ' page(s) → ' + baseName(path), 'ok');
    } catch (e) {
      console.error(e);
      setSplitStatus('Extract failed: ' + e.message, 'err');
    } finally {
      updateSplitButtons();
    }
  }

  async function splitEachPage() {
    if (!splitDoc) return;
    const pages = splitDoc.selected.size > 0
      ? [...splitDoc.selected].sort((a, b) => a - b)
      : Array.from({ length: splitDoc.pages }, (_, i) => i + 1);
    if (pages.length > 1 &&
        !confirm('This will write ' + pages.length + ' separate PDF files into your output folder. Continue?')) {
      return;
    }
    splitAllBtn.disabled = true;
    setSplitStatus('Splitting…');
    try {
      const src = await PDFLib.PDFDocument.load(splitDoc.bytes, { ignoreEncryption: true });
      for (const n of pages) {
        setSplitStatus('Splitting page ' + n + '…');
        const out = await PDFLib.PDFDocument.create();
        const [pg] = await out.copyPages(src, [n - 1]);
        out.addPage(pg);
        await saveToFolder(await out.save(), splitBaseName() + ' - p' + n + '.pdf');
      }
      setSplitStatus('Saved ' + pages.length + ' file(s) to the output folder.', 'ok');
    } catch (e) {
      console.error(e);
      setSplitStatus('Split failed: ' + e.message, 'err');
    } finally {
      updateSplitButtons();
    }
  }

  splitDrop.addEventListener('click', () => splitInput.click());
  splitInput.addEventListener('change', () => {
    if (splitInput.files[0]) loadSplitFile(splitInput.files[0]);
    splitInput.value = '';
  });
  ['dragenter', 'dragover'].forEach(ev =>
    splitDrop.addEventListener(ev, e => { e.preventDefault(); splitDrop.classList.add('drag'); }));
  ['dragleave', 'drop'].forEach(ev =>
    splitDrop.addEventListener(ev, e => { e.preventDefault(); splitDrop.classList.remove('drag'); }));
  splitDrop.addEventListener('drop', e => {
    e.stopPropagation();
    if (e.dataTransfer.files && e.dataTransfer.files[0]) loadSplitFile(e.dataTransfer.files[0]);
  });

  splitRange.addEventListener('change', applyRangeInput);
  selectAllBtn.addEventListener('click', () => {
    if (!splitDoc) return;
    splitDoc.selected = new Set(Array.from({ length: splitDoc.pages }, (_, i) => i + 1));
    splitGrid.querySelectorAll('.cell').forEach(c => c.classList.add('sel'));
    syncRangeFromSelection(); updateSplitButtons();
  });
  selectNoneBtn.addEventListener('click', () => {
    if (!splitDoc) return;
    splitDoc.selected.clear();
    splitGrid.querySelectorAll('.cell').forEach(c => c.classList.remove('sel'));
    syncRangeFromSelection(); updateSplitButtons();
  });
  extractBtn.addEventListener('click', extractSelected);
  splitAllBtn.addEventListener('click', splitEachPage);
  splitClearBtn.addEventListener('click', () => { splitDoc = null; renderSplit(); setSplitStatus(''); });

  /* ===================== Compress ===================== */
  const compressDrop = document.getElementById('compressDrop');
  const compressInput = document.getElementById('compressInput');
  const compressPanel = document.getElementById('compressPanel');
  const compressThumb = document.getElementById('compressThumb');
  const compressName = document.getElementById('compressName');
  const compressMeta = document.getElementById('compressMeta');
  const levelBtns = document.getElementById('levelBtns');
  const compressBtn = document.getElementById('compressBtn');
  const compressSaveBtn = document.getElementById('compressSaveBtn');
  const compressClearBtn = document.getElementById('compressClearBtn');
  const compressStatus = document.getElementById('compressStatus');
  const compressResult = document.getElementById('compressResult');

  // dpi = render resolution, quality = JPEG quality (0–1). Lower = smaller file.
  const LEVELS = {
    light:    { dpi: 150, quality: 0.8 },
    balanced: { dpi: 120, quality: 0.65 },
    strong:   { dpi: 96,  quality: 0.5 },
    extreme:  { dpi: 72,  quality: 0.4 },
  };
  let compressLevel = 'balanced';
  let compressDoc = null;     // { name, bytes, pages }
  let compressedBytes = null;

  function setCompressStatus(msg, cls = '') {
    compressStatus.textContent = msg;
    compressStatus.className = 'status ' + cls;
  }

  function resetCompressOutput() {
    compressedBytes = null;
    compressSaveBtn.classList.add('hidden');
    compressResult.classList.add('hidden');
  }

  async function loadCompressFile(file) {
    if (!file || !(file.type === 'application/pdf' || file.name.toLowerCase().endsWith('.pdf'))) {
      setCompressStatus('Please choose a PDF file.', 'err');
      return;
    }
    setCompressStatus('Loading…');
    resetCompressOutput();
    if (!await ready(setCompressStatus)) return;
    try {
      const bytes = new Uint8Array(await file.arrayBuffer());
      const pdf = await pdfjsLib.getDocument({ data: bytes.slice(0) }).promise;
      compressDoc = { name: file.name, bytes, pages: pdf.numPages };
      compressName.textContent = file.name;
      compressMeta.textContent = pdf.numPages + (pdf.numPages === 1 ? ' page · ' : ' pages · ') +
        fmtSize(file.size) + ' original';
      compressThumb.src = await renderPageToDataURL(pdf, 1, 0.5);
      await pdf.destroy();
      compressPanel.classList.remove('hidden');
      compressBtn.disabled = false;
      compressClearBtn.classList.remove('hidden');
      setCompressStatus('');
    } catch (e) {
      console.error(e);
      setCompressStatus('Could not read this PDF (it may be encrypted).', 'err');
      compressDoc = null;
      compressPanel.classList.add('hidden');
    }
  }

  function canvasToJpegBytes(canvas, q) {
    return new Promise(resolve => {
      canvas.toBlob(async blob => resolve(new Uint8Array(await blob.arrayBuffer())), 'image/jpeg', q);
    });
  }

  async function runCompress() {
    if (!compressDoc) return;
    const { dpi, quality } = LEVELS[compressLevel];
    compressBtn.disabled = true;
    resetCompressOutput();
    try {
      const pdf = await pdfjsLib.getDocument({ data: compressDoc.bytes.slice(0) }).promise;
      const out = await PDFLib.PDFDocument.create();
      const scale = dpi / 72;
      for (let n = 1; n <= pdf.numPages; n++) {
        setCompressStatus('Compressing page ' + n + ' / ' + pdf.numPages + '…');
        const page = await pdf.getPage(n);
        const base = page.getViewport({ scale: 1 });   // page size in points (1pt = 1/72")
        const vp = page.getViewport({ scale });
        const canvas = document.createElement('canvas');
        canvas.width = Math.max(1, Math.round(vp.width));
        canvas.height = Math.max(1, Math.round(vp.height));
        const ctx = canvas.getContext('2d');
        ctx.fillStyle = '#fff';
        ctx.fillRect(0, 0, canvas.width, canvas.height);
        await page.render({ canvasContext: ctx, viewport: vp }).promise;
        const jpg = await canvasToJpegBytes(canvas, quality);
        const img = await out.embedJpg(jpg);
        const p = out.addPage([base.width, base.height]);
        p.drawImage(img, { x: 0, y: 0, width: base.width, height: base.height });
        canvas.width = canvas.height = 0; // release memory
      }
      await pdf.destroy();
      compressedBytes = await out.save();
      const orig = compressDoc.bytes.length;
      const now = compressedBytes.length;
      const pct = Math.round((1 - now / orig) * 100);
      compressResult.classList.remove('hidden');
      if (now < orig) {
        compressResult.innerHTML = '<b>' + fmtSize(orig) + '</b> &rarr; <b>' + fmtSize(now) +
          '</b> &nbsp; <span class="save">−' + pct + '%</span>';
        setCompressStatus('Done.', 'ok');
      } else {
        compressResult.innerHTML = '<b>' + fmtSize(orig) + '</b> &rarr; <b>' + fmtSize(now) +
          '</b> &nbsp; <span class="grow">+' + Math.abs(pct) + '%</span> — this PDF is mostly text/vector, ' +
          'so image compression makes it bigger. Keep the original, or try a lighter level.';
        setCompressStatus('Done — no size saving on this file.', '');
      }
      compressSaveBtn.classList.remove('hidden');
    } catch (e) {
      console.error(e);
      setCompressStatus('Compression failed: ' + e.message, 'err');
    } finally {
      compressBtn.disabled = !compressDoc;
    }
  }

  function compressBaseName() { return compressDoc.name.replace(/\\.pdf$/i, ''); }

  compressDrop.addEventListener('click', () => compressInput.click());
  compressInput.addEventListener('change', () => {
    if (compressInput.files[0]) loadCompressFile(compressInput.files[0]);
    compressInput.value = '';
  });
  ['dragenter', 'dragover'].forEach(ev =>
    compressDrop.addEventListener(ev, e => { e.preventDefault(); compressDrop.classList.add('drag'); }));
  ['dragleave', 'drop'].forEach(ev =>
    compressDrop.addEventListener(ev, e => { e.preventDefault(); compressDrop.classList.remove('drag'); }));
  compressDrop.addEventListener('drop', e => {
    e.stopPropagation();
    if (e.dataTransfer.files && e.dataTransfer.files[0]) loadCompressFile(e.dataTransfer.files[0]);
  });
  levelBtns.addEventListener('click', e => {
    const btn = e.target.closest('.level');
    if (!btn) return;
    compressLevel = btn.dataset.level;
    levelBtns.querySelectorAll('.level').forEach(b => b.classList.toggle('active', b === btn));
    resetCompressOutput();      // a new level invalidates the previous result
    setCompressStatus('');
  });
  compressBtn.addEventListener('click', runCompress);
  compressSaveBtn.addEventListener('click', async () => {
    if (!compressedBytes) return;
    compressSaveBtn.disabled = true;
    try {
      const path = await saveToFolder(compressedBytes, compressBaseName() + ' (compressed).pdf');
      setCompressStatus('Saved ' + baseName(path), 'ok');
    } catch (e) {
      setCompressStatus('Save failed: ' + e.message, 'err');
    } finally {
      compressSaveBtn.disabled = false;
    }
  });
  compressClearBtn.addEventListener('click', () => {
    compressDoc = null;
    resetCompressOutput();
    compressPanel.classList.add('hidden');
    compressClearBtn.classList.add('hidden');
    setCompressStatus('');
  });
})();
</script>
</body>
</html>
"""


class Handler(BaseHTTPRequestHandler):
    def log_message(self, *args):
        pass

    def _json(self, code: int, payload: dict):
        body = json.dumps(payload).encode()
        self.send_response(code)
        self.send_header("Content-Type", "application/json")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def _read_json(self) -> dict:
        length = int(self.headers.get("Content-Length", 0))
        return json.loads(self.rfile.read(length))

    def do_GET(self):
        if self.path in ("/", "/index.html"):
            body = (HTML_PAGE
                    .replace("__DEFAULT_OUTPUT__", html.escape(str(DEFAULT_OUTPUT), quote=True))
                    .replace("__WORKER_PORT__", str(WORKER_PORT))
                    .replace("__VERSION__", VERSION)).encode()
            self.send_response(200)
            self.send_header("Content-Type", "text/html; charset=utf-8")
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)
        elif self.path == "/ocr/state":
            self._json(200, {"state": OCR["state"], "detail": OCR["detail"],
                             "log": list(OCR["log"])[-40:]})
        elif self.path.startswith("/word/state"):
            job_id = (parse_qs(urlparse(self.path).query).get("job") or [""])[0]
            job = WORD_JOBS.get(job_id)
            if job is None:
                self._json(200, {"ok": False, "error": "That conversion is no longer being tracked."})
                return
            self._json(200, {"ok": True, "state": job["state"], "path": job["path"],
                             "error": job["error"], "pct": _word_pct(job),
                             "note": _word_note(job)})
        elif self.path.startswith("/vendor/"):
            name = self.path[len("/vendor/"):]
            src = asset_dir() / "vendor" / name
            if name not in VENDOR_FILES or not src.is_file():
                self.send_error(404)
                return
            body = src.read_bytes()
            self.send_response(200)
            self.send_header("Content-Type", VENDOR_FILES[name])
            self.send_header("Content-Length", str(len(body)))
            self.send_header("Cache-Control", "max-age=86400")
            self.end_headers()
            self.wfile.write(body)
        else:
            self.send_error(404)

    def do_POST(self):
        if self.path == "/convert-doc":
            try:
                req = self._read_json()
                data = base64.b64decode(req["data"])
                out_dir = Path(req.get("outdir") or DEFAULT_OUTPUT).expanduser()
                out_path = convert_document(req["name"], data, out_dir)
                self._json(200, {"ok": True, "path": str(out_path)})
            except UserFacingError as e:
                # Already a finished sentence (1.4, the scan gate) — a
                # 'ValueError:' in front of it would only get in the way.
                self._json(200, {"ok": False, "error": str(e)})
            except Exception as e:
                self._json(200, {"ok": False, "error": f"{type(e).__name__}: {e}"})
        elif self.path == "/convert-word":
            try:
                req = self._read_json()
                data = base64.b64decode(req["data"])
                out_dir = Path(req.get("outdir") or DEFAULT_OUTPUT).expanduser()
                job_id = new_word_job(req["name"])
                threading.Thread(
                    target=run_word_job,
                    args=(job_id, req["name"], data, out_dir, req.get("pages", "")),
                    daemon=True).start()
                self._json(200, {"ok": True, "job": job_id})
            except Exception as e:
                self._json(200, {"ok": False, "error": f"{type(e).__name__}: {e}"})
        elif self.path == "/save-pdf":
            try:
                req = self._read_json()
                data = base64.b64decode(req["data"])
                out_dir = Path(req.get("outdir") or DEFAULT_OUTPUT).expanduser()
                out_path = save_pdf(req["name"], data, out_dir)
                self._json(200, {"ok": True, "path": str(out_path)})
            except Exception as e:
                self._json(200, {"ok": False, "error": f"{type(e).__name__}: {e}"})
        elif self.path == "/ocr/setup":
            if OCR["state"] in ("needs-setup", "error"):
                OCR.update(state="installing", detail="starting setup…")
                threading.Thread(target=setup_ocr_runtime, daemon=True).start()
            self._json(200, {"ok": True})
        elif self.path == "/open":
            try:
                req = self._read_json()
                folder = Path(req.get("outdir") or DEFAULT_OUTPUT).expanduser()
                folder.mkdir(parents=True, exist_ok=True)
                if sys.platform == "win32":
                    os.startfile(folder)  # noqa
                elif sys.platform == "darwin":
                    subprocess.Popen(["open", str(folder)])
                else:
                    subprocess.Popen(["xdg-open", str(folder)])
                self._json(200, {"ok": True})
            except Exception as e:
                self._json(200, {"ok": False, "error": str(e)})
        else:
            self.send_error(404)


def main():
    ap = argparse.ArgumentParser(description=f"{APP_NAME} {VERSION}")
    ap.add_argument("--port", type=int, default=PORT)
    ap.add_argument("--no-browser", action="store_true")
    ap.add_argument("--ocr-no-model", action="store_true",
                    help="dev: start the OCR worker without loading the model")
    args = ap.parse_args()
    if args.ocr_no_model:
        WORKER_EXTRA_ARGS.append("--no-model")

    server = None
    port = args.port
    for candidate in range(args.port, args.port + 10):  # double-launch friendly
        try:
            server = ThreadingHTTPServer(("127.0.0.1", candidate), Handler)
            port = candidate
            break
        except OSError:
            continue
    if server is None:
        print("No free port found — is MarkdownFactory already running?")
        sys.exit(1)

    threading.Thread(target=init_ocr, daemon=True).start()
    url = f"http://127.0.0.1:{port}"
    print(f"{APP_NAME} {VERSION} running at {url}  (Ctrl+C to stop)", flush=True)
    try:
        import pyi_splash  # frozen build: close the "starting…" splash
        pyi_splash.close()
    except Exception:
        pass
    if not args.no_browser:
        threading.Timer(0.5, lambda: webbrowser.open(url)).start()
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        print("\nStopped.")


if __name__ == "__main__":
    main()
